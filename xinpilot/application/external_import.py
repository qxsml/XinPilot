"""Vendor-independent transfer normalization; no GUI or chemistry dependencies.

Readers deliberately accept tables, not executable protocol languages. Every
row is a single source-to-destination transfer with an explicit volume unit.
"""
from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import math
import re
from pathlib import Path

FIELDS = ("source_labware", "source_well", "destination_labware", "destination_well", "volume", "chemical_id")
ALIASES = {
    "source_labware": ("source plate", "source labware", "source rack", "sourceplate", "source.labware"),
    "source_well": ("source well", "source position", "sourcewell", "source.well"),
    "destination_labware": ("destination plate", "destination labware", "dest plate", "target plate", "destination.labware"),
    "destination_well": ("destination well", "dest well", "target well", "destination.well"),
    "volume": ("volume", "transfer volume", "volume ul", "volume nl", "volume ml", "transfer volume ul"),
    "chemical_id": ("chemical id", "compound id", "reagent id", "id"),
}
FORMATS = {"24-well horizontal": (4, 6), "48-well horizontal": (6, 8),
           "HPLC vials vertical": (6, 8), "96-well horizontal": (8, 12),
           "96-well vertical": (8, 12), "384-well horizontal": (16, 24)}
UNITS = {"uL": 1.0, "mL": 1000.0, "nL": .001}
MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100000


def _text(value):
    return "" if value is None else str(value).strip()


def _key(value):
    return re.sub(r"[^a-z0-9]", "", _text(value).lower().replace("µ", "u").replace("μ", "u"))


def suggest_columns(headers):
    result = {}
    for field in FIELDS:
        aliases = {_key(field), *map(_key, ALIASES[field])}
        matches = [h for h in headers if _key(h) in aliases]
        result[field] = matches[0] if len(matches) == 1 else ""
    return result


def _flatten(row, prefix=""):
    result = {}
    for key, value in row.items():
        name = prefix + str(key)
        if isinstance(value, dict):
            fields = _flatten(value, name + ".")
        elif isinstance(value, list):
            raise ValueError("Nested arrays are unsupported. Export one transfer per table row.")
        else:
            fields = {name: value}
        overlap = result.keys() & fields.keys()
        if overlap:
            raise ValueError("Ambiguous JSON fields: " + ", ".join(sorted(overlap)))
        result.update(fields)
    return result


def read_table(path, sheet=None):
    """Read a flat table and retain the exact original bytes in portable state."""
    path = Path(path)
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("Import exceeds the 25 MB limit. Export a smaller transfer table.")
    raw = path.read_bytes()
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        try:
            populated = [s for s in book.worksheets if s.max_row > 1]
            if sheet:
                ws = book[sheet]
            elif len(populated) == 1:
                ws = populated[0]
            else:
                raise ValueError("Choose an Excel worksheet: " + ", ".join(book.sheetnames))
            matrix = []
            for cells in ws.iter_rows():
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("Use an Excel export with values, not formulas.")
                # Preserve the original workbook separately; keep its table
                # representation JSON-safe (dates and other Excel scalar types).
                matrix.append([_text(c.value) for c in cells])
                if len(matrix) > MAX_ROWS + 1:
                    raise ValueError("Too many rows.")
        finally:
            book.close()
        headers = [_text(v) for v in matrix[0]] if matrix else []
        rows = [dict(zip(headers, row)) for row in matrix[1:] if any(_text(v) for v in row)]
    elif suffix == ".json":
        data = json.loads(raw.decode("utf-8-sig"))
        if isinstance(data, dict):
            candidates = [v for k, v in data.items() if k.lower() in {"transfers", "worklist", "rows"} and isinstance(v, list)]
            if len(candidates) != 1:
                raise ValueError("JSON needs a transfer array or one 'transfers', 'worklist' or 'rows' array.")
            data = candidates[0]
        if not isinstance(data, list) or any(not isinstance(r, dict) for r in data):
            raise ValueError("JSON must contain transfer objects.")
        rows = [_flatten(r) for r in data]
        headers = list(dict.fromkeys(k for row in rows for k in row))
    elif suffix in {".csv", ".tsv", ".txt"}:
        text = raw.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t")
        except csv.Error as exc:
            raise ValueError("Expected a delimited table with column headers (CSV, TSV or semicolon-separated).") from exc
        reader = csv.reader(io.StringIO(text), dialect)
        headers = [_text(v) for v in next(reader, [])]
        rows = []
        for number, values in enumerate(reader, 2):
            if not any(_text(v) for v in values):
                continue
            if len(values) != len(headers):
                raise ValueError(f"Row {number}: column count differs from the header.")
            rows.append(dict(zip(headers, values)))
            if len(rows) > MAX_ROWS:
                raise ValueError("Too many rows.")
    else:
        raise ValueError("Unsupported format. Export a CSV, TSV, Excel or JSON transfer table; executable protocols are not imported.")
    if not rows or not headers or len(rows) > MAX_ROWS:
        raise ValueError("The table is empty or exceeds 100,000 transfers.")
    if any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("Column headers must be nonblank and unique.")
    return {"headers": headers, "rows": rows, "original": {
        "name": path.name, "sha256": hashlib.sha256(raw).hexdigest(),
        "base64": base64.b64encode(raw).decode("ascii")}}


def package_originals(zf, state, prefix="assets/external-import"):
    """Expose byte-exact originals as files; embedded copies also preserve history."""
    originals = [state.get("table", {}).get("original", {})] + state.get("stock_originals", [])
    for number, item in enumerate(originals, 1):
        if not item:
            continue
        raw = base64.b64decode(item["base64"], validate=True)
        if hashlib.sha256(raw).hexdigest() != item["sha256"]:
            raise ValueError("The preserved external import file failed its integrity check.")
        name = re.sub(r"[^A-Za-z0-9_.-]", "_", item["name"]).strip(".") or "original"
        zf.writestr(f"{prefix}/{number}-{name}", raw)


def positive_number(value, label):
    try:
        number = float(value)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{label}: enter a number.") from exc
    if not math.isfinite(number) or number <= 0:
        raise ValueError(f"{label}: must be finite and greater than zero.")
    return number


def well_label(value):
    text = _text(value).upper()
    match = re.fullmatch(r"([A-Z]+)0*([1-9][0-9]*)", text)
    if not match:
        raise ValueError(f"Well {text!r}: use an explicit label such as A1, not a numeric index.")
    return match[1] + str(int(match[2]))


def source_key(labware, well):
    return json.dumps([_text(labware), well_label(well)], ensure_ascii=False)


def normalize(table, mapping, unit, destination_format, destination_name=""):
    if unit not in UNITS or destination_format not in FORMATS:
        raise ValueError("Select the volume unit and destination format before previewing.")
    required = FIELDS[:2] + ("destination_well", "volume")
    for field in required:
        if mapping.get(field) not in table["headers"]:
            raise ValueError(f"Select a column for {field.replace('_', ' ')}.")
    selected = [h for h in mapping.values() if h]
    if len(selected) != len(set(selected)):
        raise ValueError("Each mapped column must have one meaning.")
    nr, nc = FORMATS[destination_format]
    transfers, plates = [], set()
    for number, row in enumerate(table["rows"], 2):
        values = {f: _text(row.get(mapping.get(f, ""))) for f in FIELDS}
        try:
            if not values["source_labware"]:
                raise ValueError("Source labware is missing.")
            values["source_well"] = well_label(values["source_well"])
            well = well_label(values["destination_well"])
            m = re.fullmatch(r"([A-Z]+)(\d+)", well)
            if len(m[1]) != 1 or ord(m[1]) - 65 >= nr or int(m[2]) > nc:
                raise ValueError(f"Destination {well} is outside {destination_format}.")
            plate = values["destination_labware"] or destination_name.strip()
            if not plate:
                raise ValueError("Destination labware is missing; map its column or enter its name.")
            plates.add(plate)
            volume_ul = positive_number(values.pop("volume"), "Volume") * UNITS[unit]
            positive_number(volume_ul, "Volume in uL")
            values.update(destination_well=well, destination_labware=plate, volume_ul=volume_ul)
            transfers.append(values)
        except ValueError as exc:
            raise ValueError(f"Row {number}: {exc}") from exc
    if len(plates) != 1:
        raise ValueError("This import supports one reaction destination plate. Found: " + ", ".join(sorted(plates)))
    return transfers


def source_rows(transfers, saved=None):
    sources = {}
    for row in transfers:
        key = source_key(row["source_labware"], row["source_well"])
        item = sources.setdefault(key, {"source_labware": row["source_labware"], "source_well": row["source_well"],
                                       "chemical_id": "", "is_building_block": False})
        incoming = row.get("chemical_id", "")
        if incoming and item["chemical_id"] and incoming != item["chemical_id"]:
            raise ValueError(f"Conflicting Chemical IDs at {row['source_labware']} / {row['source_well']}.")
        if incoming:
            item["chemical_id"] = incoming
    for key, previous in (saved or {}).items():
        if key in sources:
            # Saved mappings are explicit user choices for a physical source.
            # Choosing a new file clears them; revalidation must preserve them.
            sources[key].update(chemical_id=previous.get("chemical_id", ""), is_building_block=bool(previous.get("is_building_block")))
    return sources


def pushed_rows(transfers, sources, compounds, destination_format):
    index = {}
    for compound in compounds:
        cid = _text(compound.get("ID"))
        if cid:
            index.setdefault(cid, []).append(compound)
    result = []
    for number, row in enumerate(transfers, 1):
        stock = sources[source_key(row["source_labware"], row["source_well"])]
        cid = stock.get("chemical_id", "")
        matches = index.get(cid, [])
        if len(matches) != 1:
            raise ValueError(f"{row['source_labware']} / {row['source_well']}: Chemical ID {cid!r} must match exactly one Compound Notes record.")
        record = matches[0]
        if stock.get("is_building_block") and not _text(record.get("SMILES")):
            raise ValueError(f"Building block {cid} needs a SMILES structure in Compound Notes.")
        raw_conc = _text(record.get("Conc."))
        concentration = positive_number(raw_conc, f"{cid} concentration (M)") if raw_conc else None
        result.append({"source": "External Import", "block": number, "well": row["destination_well"],
                       "chemical": cid, "chemical_id": cid, "chemical_name": record.get("Name", ""),
                       "is_building_block": bool(stock.get("is_building_block")),
                       "volume": row["volume_ul"], "reaction_umol": row["volume_ul"] * concentration if concentration is not None else "",
                       "amount_unknown": concentration is None,
                       "dest_labware": row["destination_labware"], "dest_mode": destination_format,
                       "source_labware": row["source_labware"], "source_well": row["source_well"]})
    return result


def validate_state(state):
    """Validate optional external state before replacing a live project."""
    if not state:
        return
    for key in ("table", "sources", "settings"):
        if not isinstance(state.get(key, {}), dict):
            raise ValueError(f"External import {key} must be an object.")
    for key in ("transfers", "stock_originals"):
        if not isinstance(state.get(key, []), list):
            raise ValueError(f"External import {key} must be a list.")
    for key, source in state.get("sources", {}).items():
        if not isinstance(source, dict) or key != source_key(source.get("source_labware", ""), source.get("source_well", "")):
            raise ValueError("Invalid external source mapping.")
        if not isinstance(source.get("is_building_block", False), bool):
            raise ValueError("Building-block flags must be boolean values.")
    reviewed = state.get("reviewed_settings")
    if reviewed is not None:
        if not isinstance(reviewed, dict):
            raise ValueError("Invalid reviewed external-import settings.")
        rows = normalize(state["table"], reviewed.get("columns", {}), reviewed.get("unit"), reviewed.get("format"), reviewed.get("destination", ""))
        if rows != state.get("transfers"):
            raise ValueError("Stored external transfers differ from the reviewed original table.")
