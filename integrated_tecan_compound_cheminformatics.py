from __future__ import annotations

import csv
import itertools
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    from rdkit import Chem, rdBase
    from rdkit.Chem import AllChem, Descriptors, rdMolDescriptors
except Exception:
    Chem = None
    rdBase = None
    AllChem = None
    Descriptors = None
    rdMolDescriptors = None


APP_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
ROOT = APP_DIR.parent if APP_DIR.name.lower().startswith("dist") else APP_DIR
DEFAULT_STOCK_XLSX = ROOT / "BB Stock list - Copy (2).xlsx"
LOCAL_STOCK_XLSX = ROOT / "integrated_building_blocks.xlsx"
OUTPUT_DIR = ROOT / "outputs"

PROTON = 1.00727646688
SODIUM = 22.989218
WELL_RE = re.compile(r"^([A-Za-z]+)([0-9]+)$")
RUN_RE = re.compile(r"run\((?P<quote>['\"])(?P<smarts>.*?)(?P=quote),\s*(?P<out>['\"].*?['\"]),\s*(?P<label>['\"].*?['\"])\)")

STOCK_HEADERS = [
    "ID",
    "SMILES",
    "Name",
    "Role",
    "MW",
    "CAS",
    "Catalogue NO.",
    "Supplier",
    "Quantity (mg)",
    "Quantity(umol)",
    "Conc.",
    "Solvent",
    "Nature",
]

ROLE_OPTIONS = ["Nucleophile", "Electrophile", "Product", "Reagent / ignore"]

FALLBACK_REACTIONS = {
    "urea_isocyanate": {
        "label": "Make Urea by Isocyanate",
        "smarts": "[*:2][N:1]=[C:3]=O.[N;H1,H2:4]>>[*:2][N:1][C:3](=O)[N:4]",
        "reactant_labels": ["Isocyanate", "Amine"],
    },
    "curtius": {
        "label": "Curtius Rearrangement",
        "smarts": "[*:2][C:1](=O)[O;H1].[N;H1,H2:4]>>[*:2]N[C:1](=O)[N:4]",
        "reactant_labels": ["Carboxylic acid", "Amine"],
    },
    "acid_to_amine": {
        "label": "Acid to Amine",
        "smarts": "[*:1][C:2](=O)[O;H1]>>[*:1][NH2]",
        "reactant_labels": ["Carboxylic acid substrate"],
    },
    "isocyanate_to_amine": {
        "label": "Isocyanate to Amine",
        "smarts": "[*:1][N:2]=C=O>>[*:1][N:2]",
        "reactant_labels": ["Isocyanate substrate"],
    },
    "suzuki": {
        "label": "Suzuki Coupling",
        "smarts": "[*:1][Br,Cl,I:2].[*:3]B([O])[O]>>[*:1][*:3]",
    },
    "buchwald_hartwig": {
        "label": "Buchwald-Hartwig Amination",
        "smarts": "[c:1][Br,Cl,I:2].[N;H1,H2:3]>>[c:1][N:3]",
    },
    "amide_coupling": {
        "label": "Amide Coupling",
        "smarts": "[*:1][C:2](=O)[O;H1].[N;H1,H2:3]>>[*:1][C:2](=O)[N:3]",
    },
    "nitro_reduction": {
        "label": "Nitro Reduction to Amine",
        "smarts": "[*:1][N+:2](=[O:3])[O-:4]>>[*:1][N+0H2:2]",
        "reactant_labels": ["Nitro substrate"],
    },
    "alcohol_oxidation": {
        "label": "Alcohol Oxidation to Aldehyde/Ketone",
        "smarts": "[C;H1,H2:1][O;H1:2]>>[C:1]=[O:2]",
        "reactant_labels": ["Alcohol substrate"],
    },
    "carbonyl_reduction": {
        "label": "Carbonyl Reduction to Alcohol",
        "smarts": "[C;!$(C(=O)[!#6;!#1]):1]=[O:2]>>[C:1][O:2]",
        "reactant_labels": ["Carbonyl substrate"],
    },
    "alkene_hydrogenation": {
        "label": "Alkene Hydrogenation",
        "smarts": "[C:1]=[C:2]>>[C:1]-[C:2]",
        "reactant_labels": ["Alkene substrate"],
    },
    "ester_hydrolysis": {
        "label": "Ester Hydrolysis / Saponification",
        "smarts": "[*:1][C:2](=[O:3])[O:4][#6:5]>>[*:1][C:2](=[O:3])[OH:4]",
        "reactant_labels": ["Ester substrate"],
    },
    "amide_hydrolysis": {
        "label": "Amide Hydrolysis",
        "smarts": "[*:1][C:2](=[O:3])[N:4]>>[*:1][C:2](=[O:3])O",
        "reactant_labels": ["Amide substrate"],
    },
    "cbz_deprotection": {
        "label": "Cbz Deprotection",
        "smarts": "[N:1]C(=O)O[CH2]c1ccccc1>>[N:1]",
        "reactant_labels": ["Cbz-protected substrate"],
    },
    "silyl_ether_deprotection": {
        "label": "Silyl Ether Deprotection",
        "smarts": "[O:1][Si]([#6])([#6])[#6]>>[OH:1]",
        "reactant_labels": ["Silyl-protected substrate"],
    },
    "carbamate_isocyanate": {
        "label": "Carbamate Formation from Isocyanate",
        "smarts": "[*:1][N:2]=C=O.[O;H1:3][*:4]>>[*:1][N:2]C(=O)[O:3][*:4]",
        "reactant_labels": ["Isocyanate", "Alcohol"],
    },
    "thiourea_isothiocyanate": {
        "label": "Thiourea Formation from Isothiocyanate",
        "smarts": "[*:1][N:2]=C=S.[N;H1,H2:3]>>[*:1][N:2]C(=S)[N:3]",
        "reactant_labels": ["Isothiocyanate", "Amine"],
    },
    "alcohol_sulfonylation": {
        "label": "Alcohol Sulfonylation",
        "smarts": "[*:1]S(=O)(=O)Cl.[O;H1:2][*:3]>>[*:1]S(=O)(=O)[O:2][*:3]",
        "reactant_labels": ["Sulfonyl halide", "Alcohol"],
    },
    "n_alkylation": {
        "label": "N-Alkylation",
        "smarts": "[*:1][Cl,Br,I].[N;H1,H2:2]>>[*:1][N:2]",
        "reactant_labels": ["Alkyl halide", "Amine"],
    },
    "o_alkylation": {
        "label": "O-Alkylation",
        "smarts": "[*:1][Cl,Br,I].[O;H1:2]>>[*:1][O:2]",
        "reactant_labels": ["Alkyl halide", "Alcohol/phenol"],
    },
    "oxime_formation": {
        "label": "Oxime Formation",
        "smarts": "[C;!$(C(=O)[!#6;!#1]):1](=O).[N;H2:2][O;H1:3]>>[C:1]=[N:2][O:3]",
        "reactant_labels": ["Carbonyl compound", "Hydroxylamine"],
    },
    "amine_epoxide_opening": {
        "label": "Amine Epoxide Opening",
        "smarts": "[C:1]1[O:2][C:3]1.[N;H1,H2:4]>>[C:1]([OH:2])[C:3][N:4]",
        "reactant_labels": ["Epoxide", "Amine"],
    },
    "petasis": {
        "label": "Petasis Borono-Mannich Reaction",
        "smarts": "[*:1][CH:2]=[O:5].[N;H1,H2:3].[*:4]B(O)O>>[*:1][CH:2]([N:3])[*:4]",
        "reactant_labels": ["Aldehyde", "Amine", "Boronic acid"],
    },
}


REACTION_METADATA = {
    "urea_isocyanate": ["Isocyanate", "Amine"],
    "curtius": ["Carboxylic acid", "Amine"],
    "acid_to_amine": ["Carboxylic acid substrate"],
    "isocyanate_to_amine": ["Isocyanate substrate"],
    "suzuki": ["Aryl/vinyl halide", "Organoboron component"],
    "buchwald_hartwig": ["Aryl halide", "Amine"],
    "snar": ["Activated aryl halide", "Nucleophile"],
    "amide_coupling": ["Carboxylic acid", "Amine"],
    "sulfonamide_formation": ["Sulfonyl halide", "Amine"],
    "sonogashira": ["Aryl/vinyl halide", "Terminal alkyne"],
    "heck": ["Aryl/vinyl halide", "Alkene"],
    "esterification": ["Carboxylic acid", "Alcohol"],
    "reductive_amination": ["Carbonyl compound", "Amine"],
    "reductive_amination_aldehyde": ["Aldehyde", "Amine"],
    "reductive_amination_ketone": ["Ketone", "Amine"],
    "sn2_alkylation": ["Alkyl halide", "Nucleophile"],
    "click_triazole": ["Azide", "Alkyne"],
    "acyl_halide_amine_coupling": ["Acyl halide", "Amine"],
    "mitsunobu": ["Alcohol", "Acidic nucleophile"],
    "williamson_ether": ["Alcohol/phenol", "Alkyl halide"],
    "boc_protection": ["Amine", "Boc reagent"],
    "boc_deprotection": ["Boc-protected substrate"],
    "ullmann_cn": ["Aryl halide", "Amine"],
    "chan_lam": ["Boronic acid", "Amine"],
    "acyl_halide_alcohol_coupling": ["Acyl halide", "Alcohol"],
    "anhydride_amine_coupling": ["Anhydride", "Amine"],
    "activated_ester_amide": ["Activated ester", "Amine"],
}


def _find_chem_root() -> Path:
    bundle_text = getattr(sys, "_MEIPASS", "")
    bundle_dir = Path(bundle_text) if bundle_text else None
    candidates = [
        ROOT / "Cheminformatic_Super_Advance",
        APP_DIR / "Cheminformatic_Super_Advance",
        Path.cwd() / "Cheminformatic_Super_Advance",
    ]
    if bundle_dir:
        candidates.insert(0, bundle_dir / "Cheminformatic_Super_Advance")
    for candidate in candidates:
        if (candidate / "reaction_registry.json").exists():
            return candidate
    return candidates[0]


CHEM_ROOT = _find_chem_root()
REGISTRY_PATH = CHEM_ROOT / "reaction_registry.json"


@dataclass
class BuildingBlock:
    values: Dict[str, str]

    @property
    def id(self) -> str:
        return (self.values.get("ID") or "").strip()

    @property
    def name(self) -> str:
        return (self.values.get("Name") or "").strip()

    @property
    def smiles(self) -> str:
        return (self.values.get("SMILES") or "").strip()

    @property
    def conc(self) -> float:
        return parse_float(self.values.get("Conc."))

    @property
    def quantity_umol(self) -> float:
        return parse_float(self.values.get("Quantity(umol)"))


@dataclass
class TecanMap:
    path: Path
    role: str = "Nucleophile"
    chemicals: Dict[str, str] = field(default_factory=dict)
    volumes: Dict[str, float] = field(default_factory=dict)

    @property
    def wells(self) -> List[str]:
        return sorted(set(self.chemicals) | set(self.volumes), key=well_key)


def parse_float(value) -> float:
    try:
        txt = str(value or "").strip()
        if not txt:
            return 0.0
        return float(txt)
    except Exception:
        return 0.0


def well_key(label: str) -> Tuple[int, int, str]:
    m = WELL_RE.match((label or "").strip())
    if not m:
        return (9999, 9999, label)
    row = 0
    for ch in m.group(1).upper():
        row = row * 26 + (ord(ch) - ord("A") + 1)
    return (row, int(m.group(2)), label)


def split_cell_names(value: str) -> List[str]:
    parts = re.split(r"\s*\|\s*|[,;]\s*", value or "")
    return [p.strip() for p in parts if p.strip()]


def first_cell_name(value: str) -> str:
    names = split_cell_names(value)
    return names[0] if names else ""


def load_stock_excel(path: Path) -> List[BuildingBlock]:
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to read Excel files.")
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(1, c).value
            if v is not None and str(v).strip():
                headers[str(v).strip()] = c
        records: List[BuildingBlock] = []
        for r in range(2, ws.max_row + 1):
            vals: Dict[str, str] = {}
            nonempty = False
            for h in STOCK_HEADERS:
                c = headers.get(h)
                v = ws.cell(r, c).value if c else ""
                txt = "" if v is None else str(v).strip()
                vals[h] = txt
                nonempty = nonempty or bool(txt)
            if nonempty and (vals.get("ID") or vals.get("Name") or vals.get("SMILES")):
                records.append(BuildingBlock(vals))
        return records
    finally:
        wb.close()


def save_stock_excel(path: Path, records: Iterable[BuildingBlock]) -> None:
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to write Excel files.")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Building Blocks"
    ws.append(STOCK_HEADERS)
    for rec in records:
        ws.append([rec.values.get(h, "") for h in STOCK_HEADERS])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_tecan_map(csv_path: Path) -> TecanMap:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))

    header_idxs = []
    for i, row in enumerate(rows):
        if len(row) < 2:
            continue
        first = (row[0] or "").strip().lower()
        if first in {"row\\col", "row/col"} and any((c or "").strip().isdigit() for c in row[1:]):
            header_idxs.append(i)

    chemicals: Dict[str, str] = {}
    volumes: Dict[str, float] = {}

    for pos, hidx in enumerate(header_idxs):
        end = header_idxs[pos + 1] if pos + 1 < len(header_idxs) else len(rows)
        header = rows[hidx]
        cols = [int((c or "0").strip()) if (c or "").strip().isdigit() else None for c in header[1:]]
        body = rows[hidx + 1:end]
        has_letters = any(re.search(r"[A-Za-z]", cell or "") for row in body for cell in row[1:])

        for row in body:
            if not any((c or "").strip() for c in row):
                continue
            row_label = (row[0] or "").strip().upper()
            if not re.fullmatch(r"[A-Z]+", row_label):
                continue
            for j, cell in enumerate(row[1:]):
                col = cols[j] if j < len(cols) else None
                if not col:
                    continue
                well = f"{row_label}{col}"
                txt = (cell or "").strip()
                if not txt:
                    continue
                if has_letters:
                    chemicals[well] = txt
                else:
                    vol = parse_float(txt)
                    if vol:
                        volumes[well] = vol

    return TecanMap(path=csv_path, chemicals=chemicals, volumes=volumes)


def extract_smarts(script_path: Path) -> Optional[str]:
    if not script_path.exists():
        return None
    text = script_path.read_text(encoding="utf-8-sig", errors="ignore")
    smart_var = re.search(r"SMARTS\s*=\s*([\"'])(.*?)\1", text)
    if smart_var:
        return smart_var.group(2)
    m = RUN_RE.search(text)
    if m:
        return m.group("smarts")
    return None


def load_reactions() -> List[Dict[str, str]]:
    out: List[Dict[str, str]] = []
    if REGISTRY_PATH.exists():
        data = json.loads(REGISTRY_PATH.read_text(encoding="utf-8-sig"))
        for r in data.get("reactions", []):
            smarts = extract_smarts(CHEM_ROOT / r["working_dir"] / r["script"])
            if not smarts:
                continue
            reactant_count = 0
            if AllChem is not None:
                try:
                    with rdBase.BlockLogs():
                        rxn = AllChem.ReactionFromSmarts(smarts)
                    reactant_count = rxn.GetNumReactantTemplates() if rxn is not None else 0
                    if rxn is None or reactant_count not in {1, 2, 3}:
                        continue
                except Exception:
                    continue
            labels = list(REACTION_METADATA.get(r["id"], []))
            if len(labels) != reactant_count:
                labels = [f"Reactant {index}" for index in range(1, reactant_count + 1)]
            out.append({
                "id": r["id"], "label": r["label"], "smarts": smarts,
                "reactant_count": reactant_count, "reactant_labels": labels,
            })

    seen = {r["id"] for r in out}
    for key, value in FALLBACK_REACTIONS.items():
        if key not in seen:
            try:
                with rdBase.BlockLogs():
                    rxn = AllChem.ReactionFromSmarts(value["smarts"]) if AllChem is not None else None
                reactant_count = rxn.GetNumReactantTemplates() if rxn is not None else 0
            except Exception:
                continue
            if reactant_count not in {1, 2, 3}:
                continue
            labels = list(value.get("reactant_labels", REACTION_METADATA.get(key, [])))
            if len(labels) != reactant_count:
                labels = [f"Reactant {index}" for index in range(1, reactant_count + 1)]
            out.append({
                "id": key, "label": value["label"], "smarts": value["smarts"],
                "reactant_count": reactant_count, "reactant_labels": labels,
            })
    return out


def canonical_smiles(smiles: str) -> Tuple[object, str]:
    if Chem is None:
        raise RuntimeError("RDKit is required for product generation.")
    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        raise ValueError(f"Invalid SMILES: {smiles}")
    return mol, Chem.MolToSmiles(mol, canonical=True)


def run_reaction_ordered(smarts: str, reactants: Iterable[object]) -> List[str]:
    if AllChem is None or Chem is None:
        raise RuntimeError("RDKit is required for product generation.")
    with rdBase.BlockLogs():
        rxn = AllChem.ReactionFromSmarts(smarts)
        if rxn is None:
            return []
        reactants = tuple(reactants)
        if len(reactants) != rxn.GetNumReactantTemplates():
            return []
        products = []
        try:
            product_sets = rxn.RunReactants(reactants)
        except Exception:
            return []
    for pset in product_sets:
        if not pset:
            continue
        for p in pset:
            try:
                Chem.SanitizeMol(p)
            except Exception:
                continue
            products.append(Chem.MolToSmiles(p, canonical=True))
    return sorted(set(products))


def run_reaction_at_halide_site(
    smarts: str,
    reactants: Iterable[object],
    reactant_index: int,
    halide_atom_index: int,
) -> List[str]:
    """Return products formed by consuming one specifically labelled Cl/Br/I site.

    A temporary isotope label follows the chosen halide through RDKit. Products
    formed at another site retain the label and are discarded; products formed
    at the chosen site lose the labelled leaving group and are retained.
    """
    reactants = tuple(Chem.Mol(mol) for mol in reactants)
    if not (0 <= reactant_index < len(reactants)):
        return []
    target = reactants[reactant_index]
    if not (0 <= halide_atom_index < target.GetNumAtoms()):
        return []
    atom = target.GetAtomWithIdx(halide_atom_index)
    if atom.GetAtomicNum() not in {17, 35, 53}:
        return []
    marker = 997
    atom.SetIsotope(marker)
    products = run_reaction_ordered(smarts, reactants)
    selected = []
    for smiles in products:
        product = Chem.MolFromSmiles(smiles)
        if product is None:
            continue
        if any(product_atom.GetIsotope() == marker for product_atom in product.GetAtoms()):
            continue
        selected.append(Chem.MolToSmiles(product, canonical=True))
    return sorted(set(selected))


def reaction_template_matches(
    smarts: str,
    reactants: Iterable[object],
) -> List[List[Tuple[int, ...]]]:
    """Return the distinct template matches available on each ordered reactant."""
    if AllChem is None or Chem is None:
        raise RuntimeError("RDKit is required for reaction-site selection.")
    with rdBase.BlockLogs():
        rxn = AllChem.ReactionFromSmarts(smarts)
    if rxn is None:
        return []
    reactants = tuple(reactants)
    if len(reactants) != rxn.GetNumReactantTemplates():
        return []
    matches: List[List[Tuple[int, ...]]] = []
    for index, mol in enumerate(reactants):
        template = rxn.GetReactantTemplate(index)
        matches.append([
            tuple(int(atom_index) for atom_index in match)
            for match in mol.GetSubstructMatches(template, uniquify=True)
        ])
    return matches


def run_reaction_at_template_sites(
    smarts: str,
    reactants: Iterable[object],
    selected_matches: Iterable[Iterable[int]],
) -> List[str]:
    """Run a reaction at one exact template match on every ordered reactant.

    RDKit's ``_protected`` atom property prevents alternative template matches
    from reacting.  This works for retained reaction centres (carbonyls,
    amines, alkenes, and so on) as well as leaving groups, so site selection is
    no longer tied to halide chemistry.
    """
    reactants = tuple(Chem.Mol(mol) for mol in reactants)
    selected = [tuple(int(atom_index) for atom_index in match) for match in selected_matches]
    all_matches = reaction_template_matches(smarts, reactants)
    if len(selected) != len(reactants) or len(all_matches) != len(reactants):
        return []

    restricted = [Chem.Mol(mol) for mol in reactants]
    for reactant_index, matches in enumerate(all_matches):
        selected_match = selected[reactant_index]
        if selected_match not in matches:
            return []
        selected_atoms = set(selected_match)
        for match in matches:
            if match == selected_match:
                continue
            # An atom exclusive to an alternative match is sufficient to
            # suppress that match without blocking the selected one.
            for atom_index in set(match) - selected_atoms:
                restricted[reactant_index].GetAtomWithIdx(atom_index).SetBoolProp("_protected", True)
    return run_reaction_ordered(smarts, restricted)


def run_reaction(smarts: str, *reactants) -> List[str]:
    """Run a one-, two-, or three-component reaction, trying reactant order safely."""
    if len(reactants) == 1 and isinstance(reactants[0], (list, tuple)):
        reactants = tuple(reactants[0])
    else:
        reactants = tuple(reactants)
    products: List[str] = []
    seen_orders = set()
    for ordered in itertools.permutations(reactants):
        signature = tuple(Chem.MolToSmiles(mol, canonical=True) for mol in ordered)
        if signature in seen_orders:
            continue
        seen_orders.add(signature)
        products.extend(run_reaction_ordered(smarts, ordered))
    return sorted(set(products))


def mass_fields(smiles: str) -> Tuple[str, str, str, str, str, str]:
    mol = Chem.MolFromSmiles(smiles)
    ex = rdMolDescriptors.CalcExactMolWt(mol)
    mw = Descriptors.MolWt(mol)
    formula = rdMolDescriptors.CalcMolFormula(mol)
    return (
        f"{ex:.6f}",
        f"{mw:.3f}",
        formula,
        f"{ex + PROTON:.6f}",
        f"{ex + SODIUM:.6f}",
        f"{ex - PROTON:.6f}",
    )


def install_notebook_tab_feedback(root):
    if getattr(root, "_oneclick_tab_feedback_installed", False):
        return
    root._oneclick_tab_feedback_installed = True
    style = ttk.Style(root)
    style.map(
        "TButton",
        background=[("pressed", "#c5cbd2"), ("active", "#f8f8f8")],
        bordercolor=[("pressed", "#6f7883"), ("active", "#a8a8a8")],
        relief=[("pressed", "sunken"), ("!pressed", "flat")],
    )
    style.map(
        "TNotebook.Tab",
        background=[("active", "#e1e1e1")],
        relief=[("!pressed", "flat")],
    )

    def on_button_release(event):
        widget = event.widget
        if not isinstance(widget, ttk.Button):
            return
        try:
            if "disabled" in widget.state():
                return
            widget.state(["pressed"])
            widget.after(170, lambda w=widget: w.winfo_exists() and w.state(["!pressed"]))
        except Exception:
            pass

    root.bind_class("TButton", "<ButtonRelease-1>", on_button_release, add="+")


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Integrated Tecan + Compound Notes + Cheminformatic")
        self.geometry("1420x900")
        self.minsize(1180, 720)
        self.after(0, self._maximize_on_open)
        install_notebook_tab_feedback(self)

        self.stock_path_var = tk.StringVar(value=str(DEFAULT_STOCK_XLSX))
        self.output_path_var = tk.StringVar(value=str(OUTPUT_DIR / "integrated_products.csv"))
        self.constant_nu_var = tk.StringVar()
        self.constant_el_var = tk.StringVar()
        self.role_var = tk.StringVar(value=ROLE_OPTIONS[0])
        self.reaction_var = tk.StringVar()
        self.status_var = tk.StringVar(value="Ready")

        self.blocks: List[BuildingBlock] = []
        self.maps: List[TecanMap] = []
        self.reactions = load_reactions()
        if self.reactions:
            self.reaction_var.set(self.reactions[0]["label"])

        self.form_vars = {h: tk.StringVar() for h in STOCK_HEADERS}
        self.form_vars["Solvent"].set("DMA")

        self._build_ui()
        if DEFAULT_STOCK_XLSX.exists():
            self._load_stock(DEFAULT_STOCK_XLSX, copy_to_local=False)

    def _maximize_on_open(self):
        try:
            self.state("zoomed")
        except tk.TclError:
            try:
                self.attributes("-zoomed", True)
            except tk.TclError:
                pass

    def _build_ui(self):
        top = ttk.Frame(self, padding=8)
        top.pack(fill="x")
        ttk.Label(top, text="Stock file").pack(side="left")
        ttk.Entry(top, textvariable=self.stock_path_var, width=100).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(top, text="Load", command=self._pick_stock).pack(side="left", padx=3)
        ttk.Button(top, text="Save Local Stock", command=self._save_local_stock).pack(side="left", padx=3)

        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=8, pady=(0, 8))

        self.tab_run = ttk.Frame(nb, padding=8)
        self.tab_stock = ttk.Frame(nb, padding=8)
        nb.add(self.tab_run, text="Reaction Integration")
        nb.add(self.tab_stock, text="Building Blocks")

        self._build_run_tab()
        self._build_stock_tab()
        ttk.Label(self, textvariable=self.status_var, anchor="w", padding=(10, 4)).pack(fill="x")

    def _build_run_tab(self):
        setup = ttk.LabelFrame(self.tab_run, text="Reaction", padding=8)
        setup.pack(fill="x")
        ttk.Label(setup, text="Reaction type").grid(row=0, column=0, sticky="w")
        ttk.Combobox(
            setup,
            textvariable=self.reaction_var,
            values=[r["label"] for r in self.reactions],
            state="readonly",
            width=45,
        ).grid(row=0, column=1, sticky="w", padx=6)
        ttk.Label(setup, text="Constant nucleophile").grid(row=0, column=2, sticky="w", padx=(16, 0))
        ttk.Entry(setup, textvariable=self.constant_nu_var, width=28).grid(row=0, column=3, sticky="w", padx=6)
        ttk.Label(setup, text="Constant electrophile").grid(row=0, column=4, sticky="w", padx=(16, 0))
        ttk.Entry(setup, textvariable=self.constant_el_var, width=28).grid(row=0, column=5, sticky="w", padx=6)

        maps = ttk.LabelFrame(self.tab_run, text="Tecan chemical maps", padding=8)
        maps.pack(fill="both", expand=True, pady=8)
        controls = ttk.Frame(maps)
        controls.pack(fill="x", pady=(0, 6))
        ttk.Label(controls, text="Role for new files").pack(side="left")
        ttk.Combobox(controls, textvariable=self.role_var, values=ROLE_OPTIONS, state="readonly", width=20).pack(side="left", padx=6)
        ttk.Button(controls, text="Add Tecan CSV(s)", command=self._add_maps).pack(side="left", padx=3)
        ttk.Button(controls, text="Set Selected Role", command=self._set_selected_role).pack(side="left", padx=3)
        ttk.Button(controls, text="Remove Selected", command=self._remove_selected_map).pack(side="left", padx=3)
        ttk.Button(controls, text="Preview Merge", command=self._preview_merge).pack(side="left", padx=3)

        self.map_tree = ttk.Treeview(maps, columns=("role", "chemicals", "volumes", "file"), show="headings", height=12)
        for col, text, width in [
            ("role", "Role", 150),
            ("chemicals", "Chemical wells", 110),
            ("volumes", "Volume wells", 110),
            ("file", "File", 850),
        ]:
            self.map_tree.heading(col, text=text)
            self.map_tree.column(col, width=width, anchor="w")
        self.map_tree.pack(fill="both", expand=True)

        out = ttk.LabelFrame(self.tab_run, text="Output", padding=8)
        out.pack(fill="x")
        ttk.Entry(out, textvariable=self.output_path_var).pack(side="left", fill="x", expand=True, padx=(0, 6))
        ttk.Button(out, text="Browse", command=self._pick_output).pack(side="left", padx=3)
        ttk.Button(out, text="Generate Products + Usage", command=self._generate).pack(side="left", padx=3)

        log_frame = ttk.LabelFrame(self.tab_run, text="Log", padding=8)
        log_frame.pack(fill="both", expand=True, pady=(8, 0))
        self.log = tk.Text(log_frame, height=11, wrap="word")
        self.log.pack(fill="both", expand=True)

    def _build_stock_tab(self):
        left = ttk.Frame(self.tab_stock)
        left.pack(side="left", fill="both", expand=True)
        right = ttk.LabelFrame(self.tab_stock, text="Add / edit selected building block", padding=8)
        right.pack(side="left", fill="y", padx=(8, 0))

        search_row = ttk.Frame(left)
        search_row.pack(fill="x", pady=(0, 6))
        ttk.Button(search_row, text="Use Local Stock File", command=self._use_local_stock).pack(side="left", padx=3)

        self.stock_tree = ttk.Treeview(left, columns=("id", "name", "smiles", "qty", "conc", "role"), show="headings")
        for col, text, width in [
            ("id", "ID", 130),
            ("name", "Name", 210),
            ("smiles", "SMILES", 340),
            ("qty", "Qty umol", 100),
            ("conc", "Conc.", 80),
            ("role", "Role", 120),
        ]:
            self.stock_tree.heading(col, text=text)
            self.stock_tree.column(col, width=width, anchor="w")
        self.stock_tree.pack(fill="both", expand=True)
        self.stock_tree.bind("<<TreeviewSelect>>", self._on_stock_select)

        for i, h in enumerate(STOCK_HEADERS):
            ttk.Label(right, text=h).grid(row=i, column=0, sticky="w", pady=2)
            if h == "Role":
                ttk.Combobox(right, textvariable=self.form_vars[h], values=["", "Nucleophile", "Electrophile", "Reagent"], width=28).grid(row=i, column=1, sticky="ew", pady=2)
            else:
                ttk.Entry(right, textvariable=self.form_vars[h], width=34).grid(row=i, column=1, sticky="ew", pady=2)
        btns = ttk.Frame(right)
        btns.grid(row=len(STOCK_HEADERS), column=0, columnspan=2, sticky="ew", pady=(8, 0))
        ttk.Button(btns, text="Add New", command=self._add_stock_record).pack(side="left", padx=3)
        ttk.Button(btns, text="Update Selected", command=self._update_stock_record).pack(side="left", padx=3)
        ttk.Button(btns, text="Clear Form", command=self._clear_stock_form).pack(side="left", padx=3)

    def logw(self, text: str):
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.status_var.set(text)

    def _pick_stock(self):
        p = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx"), ("All files", "*.*")])
        if p:
            self.stock_path_var.set(p)
            self._load_stock(Path(p), copy_to_local=False)

    def _use_local_stock(self):
        self.stock_path_var.set(str(LOCAL_STOCK_XLSX))
        if LOCAL_STOCK_XLSX.exists():
            self._load_stock(LOCAL_STOCK_XLSX, copy_to_local=False)
        else:
            self._save_local_stock()

    def _load_stock(self, path: Path, copy_to_local: bool):
        self.blocks = load_stock_excel(path)
        if copy_to_local:
            save_stock_excel(LOCAL_STOCK_XLSX, self.blocks)
            self.stock_path_var.set(str(LOCAL_STOCK_XLSX))
        self._refresh_stock_tree()
        self.logw(f"Loaded {len(self.blocks)} building blocks from {path.name}")

    def _save_local_stock(self):
        save_stock_excel(LOCAL_STOCK_XLSX, self.blocks)
        self.stock_path_var.set(str(LOCAL_STOCK_XLSX))
        self.logw(f"Saved local stock file: {LOCAL_STOCK_XLSX}")

    def _refresh_stock_tree(self):
        self.stock_tree.delete(*self.stock_tree.get_children())
        for i, rec in enumerate(self.blocks):
            v = rec.values
            self.stock_tree.insert("", "end", iid=str(i), values=(rec.id, rec.name, rec.smiles, v.get("Quantity(umol)", ""), v.get("Conc.", ""), v.get("Role", "")))

    def _on_stock_select(self, _evt=None):
        sel = self.stock_tree.selection()
        if not sel:
            return
        rec = self.blocks[int(sel[0])]
        for h in STOCK_HEADERS:
            self.form_vars[h].set(rec.values.get(h, ""))

    def _form_record(self) -> BuildingBlock:
        vals = {h: self.form_vars[h].get().strip() for h in STOCK_HEADERS}
        if not vals["ID"]:
            raise ValueError("ID is required")
        if not vals["SMILES"]:
            raise ValueError("SMILES is required")
        return BuildingBlock(vals)

    def _add_stock_record(self):
        try:
            rec = self._form_record()
            if any(b.id.lower() == rec.id.lower() for b in self.blocks):
                raise ValueError(f"ID already exists: {rec.id}")
            self.blocks.append(rec)
            self._refresh_stock_tree()
            self.logw(f"Added building block {rec.id}")
        except Exception as exc:
            messagebox.showerror("Add failed", str(exc))

    def _update_stock_record(self):
        try:
            sel = self.stock_tree.selection()
            if not sel:
                raise ValueError("Select a row first")
            self.blocks[int(sel[0])] = self._form_record()
            self._refresh_stock_tree()
            self.logw("Updated selected building block")
        except Exception as exc:
            messagebox.showerror("Update failed", str(exc))

    def _clear_stock_form(self):
        for h in STOCK_HEADERS:
            self.form_vars[h].set("DMA" if h == "Solvent" else "")
        self.stock_tree.selection_remove(self.stock_tree.selection())

    def _add_maps(self):
        files = filedialog.askopenfilenames(filetypes=[("CSV", "*.csv"), ("All files", "*.*")])
        for f in files:
            try:
                m = read_tecan_map(Path(f))
                m.role = self.role_var.get()
                self.maps.append(m)
            except Exception as exc:
                messagebox.showerror("Map read failed", f"{f}\n\n{exc}")
        self._refresh_map_tree()

    def _refresh_map_tree(self):
        self.map_tree.delete(*self.map_tree.get_children())
        for i, m in enumerate(self.maps):
            self.map_tree.insert("", "end", iid=str(i), values=(m.role, len(m.chemicals), len(m.volumes), str(m.path)))

    def _set_selected_role(self):
        for iid in self.map_tree.selection():
            self.maps[int(iid)].role = self.role_var.get()
        self._refresh_map_tree()

    def _remove_selected_map(self):
        remove = {int(iid) for iid in self.map_tree.selection()}
        self.maps = [m for i, m in enumerate(self.maps) if i not in remove]
        self._refresh_map_tree()

    def _pick_output(self):
        p = filedialog.asksaveasfilename(defaultextension=".csv", initialdir=str(OUTPUT_DIR), filetypes=[("CSV", "*.csv")])
        if p:
            self.output_path_var.set(p)

    def _reaction(self) -> Dict[str, str]:
        label = self.reaction_var.get()
        return next((r for r in self.reactions if r["label"] == label), self.reactions[0])

    def _stock_indexes(self):
        by_id = {b.id.lower(): b for b in self.blocks if b.id}
        by_name = {b.name.lower(): b for b in self.blocks if b.name}
        return by_id, by_name

    def _resolve_block(self, token: str) -> Optional[BuildingBlock]:
        by_id, by_name = self._stock_indexes()
        t = (token or "").strip().lower()
        return by_id.get(t) or by_name.get(t)

    def _merged_role_maps(self):
        merged = {role: {} for role in ROLE_OPTIONS}
        volumes = {role: {} for role in ROLE_OPTIONS}
        for m in self.maps:
            for well, name in m.chemicals.items():
                if well in merged[m.role] and merged[m.role][well] != name:
                    merged[m.role][well] += " | " + name
                else:
                    merged[m.role][well] = name
            for well, vol in m.volumes.items():
                volumes[m.role][well] = volumes[m.role].get(well, 0.0) + vol
        return merged, volumes

    def _role_value_for_well(self, role_map: Dict[str, str], well: str, constant: str) -> str:
        if well in role_map:
            return first_cell_name(role_map[well])
        if constant.strip():
            return constant.strip()
        unique = sorted({first_cell_name(v) for v in role_map.values() if first_cell_name(v)})
        if len(unique) == 1:
            return unique[0]
        return ""

    def _preview_merge(self):
        merged, volumes = self._merged_role_maps()
        self.logw("Merged Tecan maps:")
        for role in ROLE_OPTIONS:
            if role == "Reagent / ignore":
                continue
            self.logw(f"  {role}: {len(merged[role])} chemical wells, {len(volumes[role])} volume wells")

    def _generate(self):
        try:
            if Chem is None:
                raise RuntimeError("RDKit is not installed in this Python environment.")
            reaction = self._reaction()
            merged, volumes = self._merged_role_maps()
            nu_map = merged["Nucleophile"]
            el_map = merged["Electrophile"]

            wells = set(nu_map) | set(el_map)
            wells |= set(volumes["Nucleophile"]) | set(volumes["Electrophile"])
            if not wells:
                raise ValueError("No nucleophile/electrophile Tecan wells loaded.")

            out_path = Path(self.output_path_var.get().strip() or OUTPUT_DIR / "integrated_products.csv")
            usage_path = out_path.with_name(out_path.stem + "_usage.csv")
            skipped_path = out_path.with_name(out_path.stem + "_skipped.csv")
            out_path.parent.mkdir(parents=True, exist_ok=True)

            product_rows = []
            usage: Dict[str, Dict[str, float]] = {}
            skipped = []

            for well in sorted(wells, key=well_key):
                nu_token = self._role_value_for_well(nu_map, well, self.constant_nu_var.get())
                el_token = self._role_value_for_well(el_map, well, self.constant_el_var.get())
                nu = self._resolve_block(nu_token)
                el = self._resolve_block(el_token)
                if not nu or not el:
                    skipped.append([well, nu_token, el_token, "missing building block details"])
                    continue
                if not nu.smiles or not el.smiles:
                    skipped.append([well, nu_token, el_token, "missing SMILES"])
                    continue

                try:
                    nu_mol, nu_smi = canonical_smiles(nu.smiles)
                    el_mol, el_smi = canonical_smiles(el.smiles)
                    products = run_reaction(reaction["smarts"], el_mol, nu_mol)
                except Exception as exc:
                    skipped.append([well, nu.id, el.id, f"reaction failed: {exc}"])
                    continue
                if not products:
                    skipped.append([well, nu.id, el.id, "no product generated"])
                    continue

                smi = products[0]
                exact, mw, formula, mh, mna, neg = mass_fields(smi)
                product_rows.append([
                    reaction["id"],
                    nu.id,
                    nu.name,
                    nu_smi,
                    el.id,
                    el.name,
                    el_smi,
                    f"P_{well}",
                    smi,
                    well,
                    exact,
                    mw,
                    formula,
                    mh,
                    mna,
                    neg,
                ])

                for role, bb in [("Nucleophile", nu), ("Electrophile", el)]:
                    vol = volumes[role].get(well, 0.0)
                    used_umol = vol * bb.conc if vol and bb.conc else 0.0
                    row = usage.setdefault(bb.id, {"volume_uL": 0.0, "used_umol": 0.0, "start_umol": bb.quantity_umol})
                    row["volume_uL"] += vol
                    row["used_umol"] += used_umol

            with out_path.open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow([
                    "reaction",
                    "nucleophile ID",
                    "nucleophile name",
                    "nucleophile SMILES",
                    "electrophile ID",
                    "electrophile name",
                    "electrophile SMILES",
                    "Product ID",
                    "product SMILES",
                    "well",
                    "exact mass",
                    "molecular weight",
                    "molecular formula",
                    "postive mass: +H",
                    "postive mass: +Na",
                    "negative mass",
                ])
                w.writerows(product_rows)

            with usage_path.open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["ID", "Name", "Role", "dispensed volume uL", "Conc.", "used umol", "starting umol", "remaining umol"])
                for bid, data in sorted(usage.items()):
                    bb = self._resolve_block(bid)
                    start = data["start_umol"]
                    used = data["used_umol"]
                    w.writerow([bid, bb.name if bb else "", bb.values.get("Role", "") if bb else "", f"{data['volume_uL']:.6g}", bb.values.get("Conc.", "") if bb else "", f"{used:.6g}", f"{start:.6g}", f"{start - used:.6g}" if start else ""])

            with skipped_path.open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["well", "nucleophile", "electrophile", "reason"])
                w.writerows(skipped)

            self.logw(f"Wrote products: {out_path}")
            self.logw(f"Wrote usage: {usage_path}")
            self.logw(f"Wrote skipped report: {skipped_path}")
            messagebox.showinfo("Done", f"Products: {len(product_rows)}\nSkipped: {len(skipped)}")
        except Exception as exc:
            messagebox.showerror("Generate failed", str(exc))
            self.logw(f"ERROR: {exc}")


if __name__ == "__main__":
    App().mainloop()
