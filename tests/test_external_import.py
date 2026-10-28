"""Synthetic transfer imports; no vendor examples or experimental data."""
import base64
import copy
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from xinpilot.application import external_import as service
from xinpilot.application.external_stocks import stage_stocks, write_template


class ExternalImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.path = self.root / "transfers.csv"
        self.path.write_text("Source plate,Source well,Destination plate,Destination well,Volume,Chemical ID\nRack1,A01,Reaction,A1,50,BB1\nRack2,A1,Reaction,A1,5,BASE\nRack1,A1,Reaction,B1,50,BB1\n", encoding="utf-8")
        self.table = service.read_table(self.path)
        self.mapping = service.suggest_columns(self.table["headers"])
        self.compounds = [{"ID": "BB1", "SMILES": "CCO", "Conc.": "0.2"}, {"ID": "BASE", "SMILES": "O", "Conc.": ""}]

    def normalized(self, table=None, **kwargs):
        return service.normalize(table or self.table, self.mapping, kwargs.get("unit", "uL"), kwargs.get("mode", "48-well horizontal"))

    def test_locations_units_and_amounts(self):
        transfers = self.normalized()
        sources = service.source_rows(transfers)
        self.assertEqual(len(sources), 2)
        sources[service.source_key("Rack1", "A1")]["is_building_block"] = True
        rows = service.pushed_rows(transfers, sources, self.compounds, "48-well horizontal")
        self.assertEqual([r["reaction_umol"] for r in rows], [10, "", 10])
        self.assertTrue(rows[1]["amount_unknown"])
        self.assertEqual(self.normalized(unit="nL")[0]["volume_ul"], .05)
        self.assertEqual(self.normalized(unit="mL")[0]["volume_ul"], 50000)

    def test_invalid_rows_fail_without_mutation(self):
        for field, value in (("Volume", "NaN"), ("Volume", "-1"), ("Destination well", "G1"), ("Source well", "1"), ("Source plate", ""), ("Destination plate", "Other")):
            table = copy.deepcopy(self.table)
            table["rows"][0][field] = value
            previous = copy.deepcopy(table)
            with self.subTest(field=field, value=value), self.assertRaises(ValueError):
                self.normalized(table)
            self.assertEqual(table, previous)
        with self.assertRaises(ValueError):
            self.normalized(unit="")
        with self.assertRaises(ValueError):
            self.normalized(mode="")

    def test_duplicate_transfers_are_retained(self):
        self.table["rows"].append(dict(self.table["rows"][0]))
        self.assertEqual(len(self.normalized()), 4)

    def test_json_flatten_collision_rejected_in_either_order(self):
        path = self.root / "ambiguous.json"
        for row in ({"source": {"well": "A1"}, "source.well": "A2"},
                    {"source.well": "A2", "source": {"well": "A1"}}):
            path.write_text(json.dumps([row]), encoding="utf-8")
            with self.subTest(row=row), self.assertRaisesRegex(ValueError, "Ambiguous JSON"):
                service.read_table(path)
        path.write_text(json.dumps([{"source": {"well": "A1", "labware": "Rack"}}]), encoding="utf-8")
        self.assertEqual(service.read_table(path)["rows"][0]["source.well"], "A1")

    def test_conflicting_source_ids_and_missing_compounds(self):
        rows = self.normalized()
        rows[2]["chemical_id"] = "OTHER"
        with self.assertRaises(ValueError):
            service.source_rows(rows)
        rows = self.normalized()
        for compounds in ([], self.compounds + [self.compounds[0]]):
            with self.assertRaises(ValueError):
                service.pushed_rows(rows, service.source_rows(rows), compounds, "48-well horizontal")

    def test_json_and_tabular_text(self):
        path = self.root / "nested.json"
        path.write_text(json.dumps({"transfers": [{"source": {"labware": "Rack1", "well": "A1"}, "destination": {"labware": "Reaction", "well": "B1"}, "volume": 3}]}))
        table = service.read_table(path)
        mapping = service.suggest_columns(table["headers"])
        self.assertEqual(service.normalize(table, mapping, "uL", "96-well horizontal")[0]["destination_well"], "B1")
        path = self.root / "transfer.txt"
        path.write_text(self.path.read_text().replace(",", "\t"))
        self.assertEqual(service.read_table(path)["rows"], self.table["rows"])
        path.write_text("a,b\n1,2,3\n")
        with self.assertRaises(ValueError):
            service.read_table(path)

    def test_stock_template_round_trip_and_atomic_conflicts(self):
        sources = service.source_rows(self.normalized())
        path = self.root / "stocks.xlsx"
        write_template(path, sources)
        table = service.read_table(path, "Source stocks")
        self.assertEqual(len(table["rows"]), 2)
        table["rows"][0].update(SMILES="OCC", Concentration=200, **{"Concentration unit": "mM", "Building block": "yes"})
        updated, additions = stage_stocks(table["rows"], sources, self.compounds)
        self.assertEqual(additions, [])
        self.assertTrue(updated[service.source_key("Rack1", "A1")]["is_building_block"])
        table["rows"][0]["SMILES"] = "CCC"
        old = copy.deepcopy(sources)
        with self.assertRaises(ValueError):
            stage_stocks(table["rows"], sources, self.compounds)
        self.assertEqual(sources, old)

    def test_original_bytes_and_integrity(self):
        state = {"table": self.table}
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            service.package_originals(archive, state)
        with zipfile.ZipFile(stream) as archive:
            self.assertEqual(archive.read(archive.namelist()[0]), self.path.read_bytes())
        self.table["original"]["base64"] = base64.b64encode(b"changed").decode()
        with zipfile.ZipFile(io.BytesIO(), "w") as archive, self.assertRaises(ValueError):
            service.package_originals(archive, state)

    def test_multi_sheet_requires_choice_and_formulas_rejected(self):
        from openpyxl import Workbook
        book = Workbook()
        for sheet in (book.active, book.create_sheet("Second")):
            sheet.append(["A", "B"])
            sheet.append([1, 2])
        path = self.root / "multi.xlsx"
        book.save(path)
        with self.assertRaisesRegex(ValueError, "Choose an Excel worksheet"):
            service.read_table(path)
        self.assertEqual(len(service.read_table(path, "Second")["rows"]), 1)
        book["Second"]["B2"] = "=1+1"
        book.save(path)
        with self.assertRaisesRegex(ValueError, "formulas"):
            service.read_table(path, "Second")

    def test_excel_metadata_dates_are_portable(self):
        from datetime import datetime
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(["Source plate", "Source well", "Destination plate", "Destination well", "Volume", "Date"])
        book.active.append(["Rack", "A1", "Plate", "A1", 5, datetime(2026, 10, 8)])
        path = self.root / "dated.xlsx"
        book.save(path)
        table = service.read_table(path)
        json.dumps(table)
        self.assertEqual(service.normalize(table, service.suggest_columns(table["headers"]), "uL", "96-well horizontal")[0]["volume_ul"], 5)


if __name__ == "__main__":
    unittest.main()
