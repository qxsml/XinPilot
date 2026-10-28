import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
from types import SimpleNamespace
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from compound_notes_frame import CompoundNotesFrame, CompoundRecord, DATA_HEADERS, DISPLAY_HEADER_LABELS
from xinpilot.application.stock_update import stock_updates, structure_key


class MasterStockTests(unittest.TestCase):
    def test_equivalent_smiles_and_preserve_blank_fields(self):
        result = stock_updates({'SMILES': 'CCO', 'Name': 'Ethanol', 'Solvent': 'water'},
                               {'SMILES': 'OCC', 'Name': 'new', 'Volume': '0', 'Solvent': '', 'Conc.': '0.2'})
        self.assertEqual(result, {'Volume': '0', 'Conc.': '0.2'})
        self.assertEqual(structure_key('[CH3:1][OH:2]'), structure_key('CO'))

    def test_different_structures_rejected(self):
        for left, right in [('CCO', 'CO'), ('F[C@H](Cl)Br', 'F[C@@H](Cl)Br'),
                            ('[13CH3]O', 'CO'), ('CN', 'C[NH3+]'), ('CC(=O)[O-].[Na+]', 'CC(=O)O')]:
            with self.subTest(left=left), self.assertRaises(ValueError):
                stock_updates({'SMILES': left}, {'SMILES': right, 'Volume': '50'})

    def test_unverifiable_identity_rejected(self):
        for smiles in ['', 'not-smiles']:
            for left, right in [(smiles, 'CCO'), ('CCO', smiles)]:
                with self.assertRaises(ValueError):
                    stock_updates({'SMILES': left}, {'SMILES': right})

    def test_cas_name_populates_only_blank_legacy_name(self):
        from test_compound_cas_lookup import CompoundCasLookupTests
        frame = CompoundCasLookupTests().make_frame()
        frame.cas_smiles_var.set('CCO'); frame.cas_name_var.set('ethanol')
        frame._apply_cas_to_form()
        self.assertEqual(frame.vars['Name'].get(), 'ethanol')
        frame.vars['Name'].set('My ethanol'); frame._apply_cas_to_form()
        self.assertEqual(frame.vars['Name'].get(), 'My ethanol')
        self.assertIn('Name', DATA_HEADERS)
        self.assertEqual(DISPLAY_HEADER_LABELS['Name'], 'Chemical name')

    def test_real_workbook_update_block_and_duplicate_rules(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'master.xlsx'
            wb = openpyxl.Workbook(); ws = wb.active; ws.append(DATA_HEADERS)
            original = {'ID': 'BB1', 'Name': 'ethanol', 'SMILES': 'CCO', 'Volume': '10', 'Solvent': 'water'}
            conflict = {'ID': 'BB2', 'Name': 'methanol', 'SMILES': 'CO', 'Volume': '20'}
            for row in [original, conflict]: ws.append([row.get(h, '') for h in DATA_HEADERS])
            ws.cell(2, len(DATA_HEADERS) + 1, '=1+2')
            wb.save(path); wb.close()
            incoming = [dict(original, SMILES='OCC', **{'Volume': '100', 'Conc.': '0.2', 'Quantity (mg)': '9.2', 'Quantity(umol)': '200', 'Solvent added (uL)': '1000', 'Solvent': ''}),
                        dict(conflict, SMILES='CC', Volume='999'),
                        {'ID': 'BB3', 'Name': 'ethanol', 'SMILES': 'CCO'},
                        {'ID': 'BB4', 'Name': 'new', 'SMILES': 'CCC'},
                        {'ID': 'BB4', 'Name': 'new', 'SMILES': 'CCC'}]
            frame = object.__new__(CompoundNotesFrame)
            frame._commit_current_form_to_record = Mock()
            frame.master_stock_excel_var = SimpleNamespace(get=lambda: str(path))
            frame.status_var = Mock()
            frame.records = [CompoundRecord(r) for r in incoming]
            with patch('compound_notes_frame.messagebox.showinfo'), patch('compound_notes_frame.messagebox.showwarning') as warning:
                frame._append_to_master_stock_list()
                self.assertIn('BB2', warning.call_args.args[1])
            wb = openpyxl.load_workbook(path); ws = wb.active
            rows = [{h: ws.cell(r, c + 1).value for c, h in enumerate(DATA_HEADERS)} for r in range(2, ws.max_row + 1)]
            self.assertEqual([r['ID'] for r in rows], ['BB1', 'BB2', 'BB4'])
            for key in ['Volume', 'Conc.', 'Quantity (mg)', 'Quantity(umol)', 'Solvent added (uL)']:
                self.assertEqual(rows[0][key], incoming[0][key])
            self.assertEqual(rows[0]['Solvent'], 'water')
            self.assertEqual(rows[1]['Volume'], '20')
            self.assertEqual(ws.cell(2, len(DATA_HEADERS) + 1).value, '=1+2')
            self.assertEqual(ws.cell(1, DATA_HEADERS.index('Name') + 1).value, 'Name')
            wb.close()
            # Load the saved workbook through the existing header-based importer.
            frame._deleted_source_rows = set()
            frame._refresh_tree = Mock()
            frame._set_associated_workbook_name = Mock()
            frame._mark_saved_snapshot = Mock()
            frame._load_excel(path)
            self.assertEqual(frame.records[0].values['Name'], 'ethanol')
