"""Synthetic import-to-export acceptance in the actual TypeScript WebView.

Only the native file chooser is stubbed; all application controls are operated
through their rendered DOM. No vendor samples or experimental files are used.
"""
import json
import sys
import time
import traceback
import zipfile
import uuid
from collections import deque
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from desktop.live_app import LiveBridge, html
from xinpilot.application.external_import import read_table, normalize, suggest_columns, source_rows
from xinpilot.application.external_stocks import write_template
from openpyxl import load_workbook
import webview

OUT = ROOT / "artifacts" / "external-import-window"
OUT.mkdir(parents=True, exist_ok=True)
transfer = OUT / "synthetic-transfers.csv"
transfer.write_text("Source plate,Source well,Destination plate,Destination well,Volume,Chemical ID\nRack1,A1,Reaction,A1,50,ARYL\nRack2,A1,Reaction,A1,100,BORON\nRack3,A1,Reaction,A1,20,BASE\n")
table = read_table(transfer)
stocks = OUT / "completed-source-stocks.xlsx"
write_template(stocks, source_rows(normalize(table, suggest_columns(table["headers"]), "uL", "48-well horizontal")))
book = load_workbook(stocks)
for row, smiles in ((2, "Clc1ccccc1"), (3, "OB(O)c1ccccc1"), (4, "O")):
    book.active.cell(row, 4, smiles)
    book.active.cell(row, 7, 200)
    book.active.cell(row, 8, "mM")
    book.active.cell(row, 9, "no" if row == 4 else "yes")
book.save(stocks)
export = OUT / ("synthetic-" + uuid.uuid4().hex[:8] + ".xinpilot.zip")
paths = deque()
bridge = LiveBridge()
window = webview.create_window("XinPilot external-import acceptance", html=html(), js_api=bridge, width=1280, height=900)
bridge._window = window
window.create_file_dialog = lambda *_args, **_kwargs: [str(paths.popleft())]
report = {}


def wait(expression, timeout=40):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        value = window.evaluate_js(expression)
        if value:
            return value
        time.sleep(.2)
    raise AssertionError("Timed out: " + expression)


def click(label):
    selector = "[...document.querySelectorAll('button')].find(b=>!b.disabled&&(b.getAttribute('aria-label')||b.textContent)===" + json.dumps(label) + ")"
    wait("Boolean(" + selector + ")")
    window.evaluate_js(selector + ".click()")
    time.sleep(.4)


def select(value):
    window.evaluate_js("(()=>{const s=[...document.querySelectorAll('select')].find(s=>[...s.options].some(o=>o.value===" + json.dumps(value) + "));if(!s)throw Error('Missing select');s.value=" + json.dumps(value) + ";s.dispatchEvent(new Event('change'));})()")
    time.sleep(.4)


def body_contains(text):
    return "document.body.textContent.includes(" + json.dumps(text) + ")"


def capture(name):
    import ctypes
    import os
    from ctypes import wintypes
    from PIL import ImageGrab
    time.sleep(.5)
    user32 = ctypes.windll.user32
    boxes = []
    callback_type = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)
    def collect(handle, _param):
        pid = wintypes.DWORD()
        user32.GetWindowThreadProcessId(handle, ctypes.byref(pid))
        if pid.value == os.getpid() and user32.IsWindowVisible(handle):
            rect = wintypes.RECT()
            user32.GetWindowRect(handle, ctypes.byref(rect))
            if rect.right - rect.left > 600:
                boxes.append((rect.left, rect.top, rect.right, rect.bottom))
        return True
    user32.EnumWindows(callback_type(collect), 0)
    if boxes:
        ImageGrab.grab(bbox=max(boxes, key=lambda b: (b[2]-b[0])*(b[3]-b[1])), all_screens=True).save(OUT / name)


def run():
    try:
        click("Synthesis Plan")
        click("External Import")
        paths.append(transfer)
        click("Import transfer file")
        wait(body_contains("3 transfer rows"))
        select("uL")
        select("48-well horizontal")
        capture("columns.png")
        click("Validate and preview transfers")
        wait(body_contains("Validated 3 transfers"))
        assert window.evaluate_js("document.querySelectorAll('.tree tbody tr').length") == 3
        paths.append(OUT / "blank-source-stocks.xlsx")
        click("Generate source-stock Excel template")
        wait(body_contains("Source-stock template saved"))
        paths.append(stocks)
        click("Import completed source-stock template")
        wait(body_contains("add 3 new Compound Notes records"))
        click("Yes")
        wait(body_contains("Source-stock mappings applied"))
        capture("source-stocks.png")
        click("3. Transfer preview")
        wait("document.querySelectorAll('.tree tbody tr').length===3")
        click("Push All to Cheminformatics")
        wait(body_contains("Pushed 3 transfers"))
        click("Push All to Cheminformatics")
        wait(body_contains("Pushed 3 transfers"))
        click("Reaction Summary")
        wait(body_contains("3 pushed addition(s)"))
        wait(body_contains("48-well horizontal"))
        wait(body_contains("ARYL"))
        paths.append(export)
        click("Export Project")
        deadline = time.monotonic() + 40
        while time.monotonic() < deadline:
            prompt = window.evaluate_js("document.querySelector('.prompt-box')?.textContent||''")
            if prompt and window.evaluate_js("[...document.querySelectorAll('.prompt-box button')].some(b=>b.textContent==='OK')"):
                click("OK")
            if export.exists() and not prompt:
                break
            time.sleep(.2)
        assert export.exists()
        with zipfile.ZipFile(export) as archive:
            state = json.loads(archive.read("project.json"))
            assert len(state["external_import"]["transfers"]) == 3
            assert len(state["compound_notes"]["records"]) == 3
            assert archive.read("assets/external-import/1-synthetic-transfers.csv") == transfer.read_bytes()
        info = window.evaluate_js("({unsupported:[...document.querySelectorAll('.unsupported')].map(n=>n.textContent),error:document.querySelector('#notice').textContent})")
        assert not info["unsupported"] and not info["error"], info
        report.update(status="PASS", transfer_count=3, compounds_added=3, duplicate_push=False,
                      destination="48-well horizontal", original_file_preserved=True)
    except Exception:
        report.update(status="FAIL", error=traceback.format_exc())
    finally:
        (OUT / "validation.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
        window.evaluate_js("ready=false")
        window.destroy()
        bridge._shutdown()


webview.start(run, gui="edgechromium")
print(json.dumps(report))
sys.exit(report.get("status") != "PASS")
