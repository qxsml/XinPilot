"""Synthetic end-to-end checks using the real application controllers."""
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from integrated_full_software import IntegratedFullApp
from xinpilot.application.external_import import source_key
from xinpilot.application.project_archive import save_project_package, load_project_package
from xinpilot.application.heatmap_layout import reaction_mode, heatmap_layout


class ExternalWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = IntegratedFullApp()
        cls.app.withdraw()
        cls.app.update()
        cls.original = cls.app._project_payload()

    @classmethod
    def tearDownClass(cls):
        for token in cls.app.tk.splitlist(cls.app.tk.call("after", "info")):
            cls.app.tk.call("after", "cancel", token)
        cls.app.destroy()

    def setUp(self):
        self.app._apply_project_data(self.original)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.frame = self.app.external_import_tab
        self.path = self.root / "synthetic.csv"
        self.path.write_text("Source plate,Source well,Destination plate,Destination well,Volume,Chemical ID\nRack1,A1,Reaction,A1,50,ARYL\nRack2,A1,Reaction,A1,100,BORON\nRack3,A1,Reaction,A1,20,BASE\n")
        self.app._mark_project_saved()
        with patch("external_import_frame.filedialog.askopenfilename", return_value=str(self.path)):
            self.frame._choose_file()
        self.frame.unit_var.set("uL")
        self.frame.format_var.set("48-well horizontal")
        self.frame._preview()

    def add_compounds(self):
        self.app._add_external_compounds([
            {"ID": "ARYL", "SMILES": "Clc1ccccc1", "Conc.": "0.2"},
            {"ID": "BORON", "SMILES": "OB(O)c1ccccc1", "Conc.": "0.2"},
            {"ID": "BASE", "SMILES": "O", "Conc.": ""},
        ])
        for rack in ("Rack1", "Rack2"):
            self.frame.sources[source_key(rack, "A1")]["is_building_block"] = True

    def test_push_chemistry_summary_and_portable_roundtrip(self):
        self.add_compounds()
        self.assertTrue(self.app._check_project_changes())
        chem = self.app.cheminformatics_tab
        chem.push_from_manual([{"source": "Manual workflow", "chemical_id": "BASE", "chemical": "BASE", "well": "A1", "volume": 1, "dest_mode": "48-well horizontal", "is_building_block": False}])
        self.frame._push()
        self.frame._push()
        self.assertEqual(len(chem.pushed_rows), 4)
        self.assertEqual(self.app._pushed_plan_status(), "")
        reaction = next(r for r in chem.reactions if "suzuki" in (r["id"] + r["label"]).lower())
        chem.reaction_var.set(reaction["label"])
        with patch("cheminformatics_frame.messagebox.showinfo"), patch("cheminformatics_frame.messagebox.askyesno", return_value=False):
            chem._generate()
        products = chem._pyparse_rows()
        self.assertEqual(len(products), 1)
        from rdkit import Chem
        self.assertEqual(Chem.MolToSmiles(Chem.MolFromSmiles(products[0]["desired product smiles"])), "c1ccc(-c2ccccc2)cc1")
        rows = list(chem.pushed_rows.values())
        mode, _, wells, groups = self.app.reaction_summary_tab._destination_summary([], rows)
        self.assertEqual(mode, "48-well horizontal")
        self.assertEqual(set(wells["A1"]["chemicals"]), {"ARYL", "BORON", "BASE"})
        base = next(row for values in groups.values() for row in values if row["chemical"] == "BASE" and row["gwl"] == "External Import")
        self.assertIsNone(base["umol"])
        self.assertEqual(len(heatmap_layout(reaction_mode(rows))[0]), 6)
        destination = self.root / "synthetic.xinpilot.zip"
        before = self.frame.export_state()
        save_project_package(destination, self.app._project_payload())
        with zipfile.ZipFile(destination) as archive:
            self.assertEqual(archive.read("assets/external-import/1-synthetic.csv"), self.path.read_bytes())
        restored, _ = load_project_package(destination, self.root / "reopened", normalize_well=str)
        self.app._apply_project_data(restored)
        self.assertEqual(self.frame.export_state(), before)
        self.assertEqual(len(chem.pushed_rows), 4)
        self.frame._push()
        self.assertEqual(len(chem.pushed_rows), 4)
        self.frame.unit_var.set("nL")
        with self.assertRaisesRegex(ValueError, "preview"):
            self.frame.rows_for_cheminformatics()
        self.assertIn("External Import", self.app._pushed_plan_status())

    def test_template_addition_and_mapping_reuse(self):
        from openpyxl import load_workbook
        template = self.root / "source-stocks.xlsx"
        with patch("external_import_frame.filedialog.asksaveasfilename", return_value=str(template)):
            self.frame._template()
        book = load_workbook(template)
        for row, smiles in ((2, "Clc1ccccc1"), (3, "OB(O)c1ccccc1"), (4, "O")):
            book.active.cell(row, 4, smiles)
            book.active.cell(row, 7, 200)
            book.active.cell(row, 8, "mM")
            book.active.cell(row, 9, "no" if row == 4 else "yes")
        book.save(template)
        with patch("external_import_frame.filedialog.askopenfilename", return_value=str(template)), patch("external_import_frame.messagebox.askyesno", return_value=True):
            self.frame._import_stocks()
        self.assertEqual(len(self.frame.rows_for_cheminformatics()), 3)
        self.assertEqual(len(self.frame.stock_originals), 1)
        mapping = self.root / "mapping.json"
        with patch("external_import_frame.filedialog.asksaveasfilename", return_value=str(mapping)):
            self.frame._save_mappings()
        self.frame.sources = {}
        with patch("external_import_frame.filedialog.askopenfilename", return_value=str(mapping)):
            self.frame._load_mappings()
        with self.assertRaises(ValueError):
            self.frame.rows_for_cheminformatics()
        self.frame._preview()
        self.assertTrue(self.frame.rows_for_cheminformatics()[0]["is_building_block"])

    def test_old_project_clears_external_state(self):
        self.app._apply_project_data({"version": 1})
        self.assertEqual(self.frame.transfers, [])
        self.assertEqual(self.frame.sources, {})
        self.assertEqual(self.frame.table, {})

    def test_corrected_source_survives_mapping_reload_and_preview(self):
        self.add_compounds()
        key = source_key("Rack2", "A1")
        self.frame.sources[key].update(chemical_id="CORRECTED", is_building_block=True)
        mapping = self.root / "corrected.json"
        with patch("external_import_frame.filedialog.asksaveasfilename", return_value=str(mapping)):
            self.frame._save_mappings()
        self.frame.sources = {}
        with patch("external_import_frame.filedialog.askopenfilename", return_value=str(mapping)):
            self.frame._load_mappings()
        self.frame._preview()
        self.frame._preview()
        self.assertEqual(self.frame.sources[key]["chemical_id"], "CORRECTED")
        self.assertTrue(self.frame.sources[key]["is_building_block"])
        # A genuinely new import must still start from its own file identities.
        with patch("external_import_frame.filedialog.askopenfilename", return_value=str(self.path)), patch("external_import_frame.messagebox.askyesno", return_value=True):
            self.frame._choose_file()
        self.frame.unit_var.set("uL")
        self.frame.format_var.set("48-well horizontal")
        self.frame._preview()
        self.assertEqual(self.frame.sources[key]["chemical_id"], "BORON")

    def test_conflicting_formats_rejected_by_every_route_after_external(self):
        self.add_compounds()
        self.frame._push()
        chem = self.app.cheminformatics_tab
        before = dict(chem.pushed_rows)
        for push in (chem.push_from_manual, chem.push_from_tecan, chem.push_from_opentrons):
            with self.subTest(route=push.__name__), self.assertRaisesRegex(ValueError, "conflicting"):
                push([{"chemical": "BASE", "well": "A1", "volume": 1, "dest_mode": "96-well horizontal"}])
            self.assertEqual(chem.pushed_rows, before)

    def test_changed_building_block_invalidates_custom_product_and_roundtrip(self):
        self.add_compounds()
        self.app._add_external_compounds([{"ID": "BORON2", "SMILES": "Cc1ccc(B(O)O)cc1", "Conc.": "0.2"}])
        self.frame._push()
        chem = self.app.cheminformatics_tab
        chem.custom_template_text.delete("1.0", "end")
        chem.custom_template_text.insert("1.0", "[c:1]Cl.[c:2]B(O)O>>[c:1][c:2]")
        with patch("cheminformatics_frame.messagebox.showinfo"), patch("cheminformatics_frame.messagebox.showerror") as error:
            chem._generate_custom_products(False)
        error.assert_not_called()
        original_product = dict(chem.custom_products["A1"])
        self.frame._push()
        self.assertEqual(chem.custom_products["A1"], original_product)
        # Unrelated explicit user corrections must not be deleted.
        chem.imported_product_details["B1"] = {"desired product smiles": "CCO"}
        self.frame.sources[source_key("Rack2", "A1")]["chemical_id"] = "BORON2"
        self.frame._push()
        self.assertNotIn("A1", chem.custom_products)
        self.assertEqual(chem.imported_product_details["B1"]["desired product smiles"], "CCO")
        self.assertFalse(any(row.get("desired product smiles") == original_product["smiles"] for row in chem._pyparse_rows()))
        destination = self.root / "updated.xinpilot.zip"
        save_project_package(destination, self.app._project_payload())
        restored, _ = load_project_package(destination, self.root / "reopened", normalize_well=str)
        self.app._apply_project_data(restored)
        self.assertNotIn("A1", chem.custom_products)

    def test_conflicting_destination_format_does_not_replace_records(self):
        self.add_compounds()
        chem = self.app.cheminformatics_tab
        chem.push_from_manual([{"source": "Manual workflow", "chemical": "BASE", "well": "A1", "volume": 1, "dest_mode": "96-well horizontal", "is_building_block": False}])
        before = dict(chem.pushed_rows)
        with self.assertRaisesRegex(ValueError, "conflicting"):
            self.frame._push()
        self.assertEqual(chem.pushed_rows, before)


if __name__ == "__main__":
    unittest.main()
