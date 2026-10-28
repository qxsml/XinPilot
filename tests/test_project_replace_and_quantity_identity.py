"""Protect unsaved projects and workbook chemical identity during updates."""
import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from compound_notes_frame import CompoundNotesFrame, CompoundRecord
from integrated_full_software import IntegratedFullApp


class ProjectReplacementTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = IntegratedFullApp()
        cls.app.withdraw()
        cls.app.update()
        cls.original = copy.deepcopy(cls.app._project_payload())

    @classmethod
    def tearDownClass(cls):
        for token in cls.app.tk.splitlist(cls.app.tk.call('after', 'info')):
            cls.app.tk.call('after', 'cancel', token)
        cls.app.destroy()

    def setUp(self):
        self.app._apply_project_data(copy.deepcopy(self.original))
        self.app._mark_project_saved()

    def import_with_choice(self, choice):
        with patch('integrated_full_software.filedialog.askopenfilename', return_value='incoming.xinpilot.zip'), \
             patch.object(self.app, '_load_project_package', return_value=(copy.deepcopy(self.original), None)), \
             patch('integrated_full_software.messagebox.askyesnocancel', return_value=choice) as prompt, \
             patch('integrated_full_software.messagebox.showinfo'), \
             patch('integrated_full_software.messagebox.showerror') as error:
            self.app._import_project()
        error.assert_not_called()
        return prompt

    def dirty_plan(self):
        self.app.reaction_summary_tab.temperature_var.set('130')
        self.app.manual_tab.rows.append({'reaction_well': 'A1', 'chemical_id': 'TEST', 'transfer_ul': '10'})

    def test_clean_import_does_not_prompt(self):
        self.import_with_choice(None).assert_not_called()
        self.assertFalse(self.app._check_project_changes())

    def test_cancel_preserves_unsaved_plan(self):
        self.dirty_plan()
        self.import_with_choice(None).assert_called_once()
        self.assertEqual(len(self.app.manual_tab.rows), 1)
        self.assertEqual(self.app.reaction_summary_tab.temperature_var.get(), '130')
        self.assertTrue(self.app._check_project_changes())

    def test_discard_imports_and_sets_clean_baseline(self):
        self.dirty_plan()
        self.import_with_choice(False).assert_called_once()
        self.assertEqual(self.app.manual_tab.rows, [])
        self.assertFalse(self.app._check_project_changes())

    def test_failed_or_cancelled_export_keeps_current_project(self):
        self.dirty_plan()
        with patch.object(self.app, '_export_project', return_value=False) as export:
            self.import_with_choice(True)
        export.assert_called_once()
        self.assertEqual(len(self.app.manual_tab.rows), 1)
        self.assertTrue(self.app._check_project_changes())

    def test_export_saves_old_plan_before_replacement(self):
        self.dirty_plan()
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / 'previous.xinpilot.zip'
            with patch('integrated_full_software.filedialog.asksaveasfilename', return_value=str(target)):
                self.import_with_choice(True)
            from zipfile import ZipFile
            import json
            with ZipFile(target) as archive:
                saved = json.loads(archive.read('project.json'))
            self.assertEqual(len(saved['manual_workflow']['rows']), 1)
            self.assertEqual(self.app.manual_tab.rows, [])


class QuantityIdentityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'stocks.xlsx'
        self.frame = object.__new__(CompoundNotesFrame)
        self.frame.records = [CompoundRecord({'ID': 'A', 'SMILES': 'CCO', 'Quantity (mg)': '15'}, source_row=2),
                              CompoundRecord({'ID': 'B', 'SMILES': 'CCN', 'Quantity (mg)': '20'}, source_row=3)]

    def workbook(self, rows):
        book = Workbook()
        sheet = book.active
        sheet.append(['ID', 'SMILES', 'Quantity (mg)', 'Unrelated'])
        for row in rows:
            sheet.append(row)
        book.save(self.path)
        book.close()

    def test_reordered_rows_resolve_by_id_and_keep_formulas(self):
        self.workbook([['B', 'CCN', 20, '=1+2'], ['A', 'OCC', 10, '=2+3']])
        self.frame._write_quantity_column_to_excel(self.path)
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        self.assertEqual(book.active['C2'].value, '20')
        self.assertEqual(book.active['C3'].value, '15')
        self.assertEqual(book.active['D2'].value, '=1+2')
        self.assertEqual([r.source_row for r in self.frame.records], [3, 2])

    def test_missing_duplicate_or_mismatched_identity_keeps_file_unchanged(self):
        cases = [(['A', 'CCO', 10], ['A', 'CCN', 20]),
                 (['A', 'CCO', 10], ['OTHER', 'CCN', 20]),
                 (['A', 'CCO', 10], ['B', 'CCC', 20])]
        for rows in cases:
            with self.subTest(rows=rows):
                self.workbook(rows)
                before = self.path.read_bytes()
                with self.assertRaises(ValueError):
                    self.frame._write_quantity_column_to_excel(self.path)
                self.assertEqual(self.path.read_bytes(), before)

    def test_new_unsaved_row_does_not_write_to_an_unrelated_row(self):
        self.workbook([['A', 'CCO', 10], ['B', 'CCN', 20]])
        self.frame.records.append(CompoundRecord({'ID': 'NEW', 'SMILES': 'CCC', 'Quantity (mg)': '999'}))
        self.frame._write_quantity_column_to_excel(self.path)
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        self.assertEqual(book.active.max_row, 3)
        self.assertEqual(book.active['C2'].value, '15')

    def test_one_sided_missing_structure_keeps_workbook_unchanged(self):
        for saved, current in [('CCO', ''), ('', 'CCN')]:
            with self.subTest(saved=saved, current=current):
                self.workbook([['A', saved, 10], ['B', 'CCN', 20]])
                self.frame.records[0].values['SMILES'] = current
                before = self.path.read_bytes()
                with self.assertRaisesRegex(ValueError, 'structure information'):
                    self.frame._write_quantity_column_to_excel(self.path)
                self.assertEqual(self.path.read_bytes(), before)

    def test_both_blank_structures_preserve_id_only_legacy_support(self):
        self.workbook([['A', '', 10], ['B', 'CCN', 20]])
        self.frame.records[0].values['SMILES'] = ''
        self.frame._write_quantity_column_to_excel(self.path)
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        self.assertEqual(book.active['C2'].value, '15')


if __name__ == '__main__':
    unittest.main()
