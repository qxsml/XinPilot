"""Core reaction-library enumeration for the standalone bulk GUI.

The engine deliberately keeps reaction enumeration separate from docking
preparation.  Products are canonicalized, but they are not protonated,
tautomerized, minimized, or converted to 3D here.
"""

from __future__ import annotations

import csv
import itertools
import json
import threading
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from rdkit import Chem, rdBase
from rdkit.Chem import Descriptors, rdMolDescriptors
from rdkit.Chem.rdChemReactions import ChemicalReaction, ReactionFromSmarts


SUPPORTED_INPUTS = {".csv", ".xlsx", ".xlsm", ".sdf", ".sd"}
SUPPORTED_OUTPUTS = {".csv", ".xlsx", ".sdf"}


@dataclass(frozen=True)
class ReactionTemplate:
    id: str
    label: str
    smarts: str
    reactant_labels: tuple[str, ...]
    reaction: ChemicalReaction = field(repr=False, compare=False)

    @property
    def reactant_count(self) -> int:
        return len(self.reactant_labels)


@dataclass
class BuildingBlock:
    identifier: str
    smiles: str
    properties: dict[str, str]
    mol: Chem.Mol = field(repr=False)
    source_row: int | None = None


@dataclass(frozen=True)
class SourceInspection:
    path: Path
    kind: str
    fields: tuple[str, ...]
    default_id_field: str
    default_smiles_field: str | None


@dataclass
class LoadResult:
    records: list[BuildingBlock]
    issues: list[dict[str, Any]]


@dataclass
class EnumerationResult:
    products: list[dict[str, Any]]
    failures: list[dict[str, Any]]
    total_combinations: int
    processed_combinations: int
    cancelled: bool = False

    @property
    def unique_product_count(self) -> int:
        return len({row.get("Product_SMILES", "") for row in self.products})


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _dedupe_headers(headers: Iterable[Any]) -> list[str]:
    result: list[str] = []
    counts: dict[str, int] = {}
    for index, raw in enumerate(headers, 1):
        base = _clean(raw) or f"Column_{index}"
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return result


def _best_field(fields: Sequence[str], candidates: Sequence[str]) -> str:
    lookup = {field.casefold(): field for field in fields}
    for name in candidates:
        if name.casefold() in lookup:
            return lookup[name.casefold()]
    return fields[0] if fields else ""


def load_reaction_templates(path: str | Path | None = None) -> list[ReactionTemplate]:
    template_path = Path(path) if path else Path(__file__).with_name("reaction_templates.json")
    payload = json.loads(template_path.read_text(encoding="utf-8"))
    rows = payload.get("reactions", payload) if isinstance(payload, dict) else payload
    if not isinstance(rows, list) or not rows:
        raise ValueError("Reaction template file contains no reactions")

    templates: list[ReactionTemplate] = []
    seen: set[str] = set()
    for row in rows:
        rid = _clean(row.get("id"))
        labels = tuple(_clean(value) for value in row.get("reactant_labels", []))
        smarts = _clean(row.get("smarts"))
        if not rid or rid in seen:
            raise ValueError(f"Missing or duplicate reaction ID: {rid!r}")
        with rdBase.BlockLogs():
            reaction = ReactionFromSmarts(smarts)
            if reaction is None:
                raise ValueError(f"Invalid reaction SMARTS for {rid}")
            reaction.Initialize()
        if reaction.GetNumReactantTemplates() != len(labels):
            raise ValueError(
                f"{rid} has {reaction.GetNumReactantTemplates()} SMARTS reactants "
                f"but {len(labels)} role labels"
            )
        seen.add(rid)
        templates.append(
            ReactionTemplate(rid, _clean(row.get("label")) or rid, smarts, labels, reaction)
        )
    return templates


def inspect_source(path: str | Path) -> SourceInspection:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_INPUTS:
        raise ValueError(f"Unsupported input format: {source.suffix or '(none)'}")
    if not source.is_file():
        raise FileNotFoundError(source)

    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            try:
                fields = _dedupe_headers(next(reader))
            except StopIteration:
                raise ValueError("CSV file is empty") from None
        kind = "CSV"
    elif suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            fields = _dedupe_headers(first or [])
        finally:
            workbook.close()
        if not fields:
            raise ValueError("Excel file has no header row")
        kind = "Excel"
    else:
        fields = ["_Name"]
        supplier = Chem.SDMolSupplier(str(source), removeHs=False, sanitize=False)
        prop_names: set[str] = set()
        valid_seen = False
        for molecule_number, mol in enumerate(supplier, 1):
            if mol is None:
                continue
            valid_seen = True
            prop_names.update(mol.GetPropNames(includePrivate=False, includeComputed=False))
            # Property discovery is only for configuring the GUI. Sampling keeps
            # browsing a multi-million-record SDF responsive; the full file is
            # still read during generation in the background.
            if molecule_number >= 100 or len(prop_names) >= 250:
                break
        if not valid_seen:
            raise ValueError("SDF contains no readable molecules")
        fields.extend(sorted(name for name in prop_names if name != "_Name"))
        kind = "SDF"

    id_field = _best_field(fields, ("ID", "Compound_ID", "Identifier", "Name", "_Name"))
    smiles_field = None if kind == "SDF" else _best_field(
        fields, ("SMILES", "Canonical_SMILES", "Smiles", "Structure")
    )
    return SourceInspection(source, kind, tuple(fields), id_field, smiles_field)


def _mol_to_smiles(mol: Chem.Mol) -> str:
    return Chem.MolToSmiles(mol, canonical=True, isomericSmiles=True)


def _make_block(
    identifier: Any,
    smiles: Any,
    properties: dict[str, Any],
    source_row: int | None,
) -> BuildingBlock:
    clean_id = _clean(identifier)
    clean_smiles = _clean(smiles)
    if not clean_id:
        raise ValueError("missing ID")
    if not clean_smiles:
        raise ValueError("missing SMILES")
    with rdBase.BlockLogs():
        mol = Chem.MolFromSmiles(clean_smiles)
    if mol is None:
        raise ValueError("invalid SMILES")
    canonical = _mol_to_smiles(mol)
    return BuildingBlock(
        clean_id,
        canonical,
        {str(key): _clean(value) for key, value in properties.items()},
        mol,
        source_row,
    )


def single_building_block(identifier: str, smiles: str) -> BuildingBlock:
    return _make_block(identifier, smiles, {"ID": identifier, "SMILES": smiles}, 1)


def load_building_blocks(
    path: str | Path,
    id_field: str,
    smiles_field: str | None = None,
) -> LoadResult:
    source = Path(path)
    inspection = inspect_source(source)
    records: list[BuildingBlock] = []
    issues: list[dict[str, Any]] = []

    def accept(row: dict[str, Any], row_number: int, structure_smiles: str | None = None) -> None:
        try:
            raw_smiles = structure_smiles if structure_smiles is not None else row.get(smiles_field or "")
            records.append(_make_block(row.get(id_field), raw_smiles, row, row_number))
        except Exception as exc:
            issues.append(
                {
                    "Failure_Type": "Input",
                    "Source_File": str(source),
                    "Source_Row": row_number,
                    "Input_ID": _clean(row.get(id_field)),
                    "Reason": str(exc),
                }
            )

    if inspection.kind == "CSV":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            for row_number, row in enumerate(reader, 2):
                accept(dict(row), row_number)
    elif inspection.kind == "Excel":
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, None)
            headers = _dedupe_headers(first or [])
            for row_number, values in enumerate(iterator, 2):
                row = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
                if any(_clean(value) for value in row.values()):
                    accept(row, row_number)
        finally:
            workbook.close()
    else:
        supplier = Chem.SDMolSupplier(str(source), removeHs=False, sanitize=True)
        for row_number, mol in enumerate(supplier, 1):
            if mol is None:
                issues.append(
                    {
                        "Failure_Type": "Input",
                        "Source_File": str(source),
                        "Source_Row": row_number,
                        "Input_ID": "",
                        "Reason": "unreadable or invalid SDF molecule",
                    }
                )
                continue
            row = {name: mol.GetProp(name) for name in mol.GetPropNames(False, False)}
            row["_Name"] = mol.GetProp("_Name") if mol.HasProp("_Name") else ""
            accept(row, row_number, _mol_to_smiles(mol))
    return LoadResult(records, issues)


def compatible_blocks(template: ReactionTemplate, role_index: int, blocks: Sequence[BuildingBlock]) -> tuple[list[BuildingBlock], list[dict[str, Any]]]:
    query = template.reaction.GetReactantTemplate(role_index)
    compatible: list[BuildingBlock] = []
    issues: list[dict[str, Any]] = []
    role = template.reactant_labels[role_index]
    for block in blocks:
        if block.mol.HasSubstructMatch(query):
            compatible.append(block)
        else:
            issues.append(
                {
                    "Failure_Type": "Compatibility",
                    "Role": role,
                    "Input_ID": block.identifier,
                    "Input_SMILES": block.smiles,
                    "Source_Row": block.source_row or "",
                    "Reason": f"does not match the {role} reactant pattern",
                }
            )
    return compatible, issues


def _prefix_properties(row: dict[str, Any], role_number: int, block: BuildingBlock) -> None:
    prefix = f"BB{role_number}_"
    row[prefix + "ID"] = block.identifier
    row[prefix + "SMILES"] = block.smiles
    for key, value in block.properties.items():
        safe_key = _clean(key).replace("\n", " ").replace("\r", " ")
        output_key = prefix + safe_key
        if output_key not in row:
            row[output_key] = value


def _failure_for_combination(template: ReactionTemplate, combo_id: str, blocks: Sequence[BuildingBlock], reason: str) -> dict[str, Any]:
    row: dict[str, Any] = {
        "Failure_Type": "Reaction",
        "Reaction_ID": template.id,
        "Reaction_Type": template.label,
        "Combination_ID": combo_id,
        "Reason": reason,
    }
    for index, block in enumerate(blocks, 1):
        row[f"BB{index}_ID"] = block.identifier
        row[f"BB{index}_SMILES"] = block.smiles
    return row


def enumerate_library(
    template: ReactionTemplate,
    role_records: Sequence[Sequence[BuildingBlock]],
    *,
    product_prefix: str = "PROD-",
    start_number: int = 1,
    prior_failures: Sequence[dict[str, Any]] | None = None,
    cancel_event: threading.Event | None = None,
    progress: Callable[[int, int, int], None] | None = None,
) -> EnumerationResult:
    if len(role_records) != template.reactant_count:
        raise ValueError(f"{template.label} requires {template.reactant_count} reactant roles")
    if start_number < 0:
        raise ValueError("Start number cannot be negative")

    failures = list(prior_failures or [])
    compatible_roles: list[list[BuildingBlock]] = []
    for index, blocks in enumerate(role_records):
        compatible, mismatches = compatible_blocks(template, index, blocks)
        compatible_roles.append(compatible)
        failures.extend(mismatches)

    total = 1
    for blocks in compatible_roles:
        total *= len(blocks)
    products: list[dict[str, Any]] = []
    first_product_by_smiles: dict[str, str] = {}
    processed = 0
    next_product_number = start_number

    for combination_number, blocks_tuple in enumerate(itertools.product(*compatible_roles), 1):
        if cancel_event is not None and cancel_event.is_set():
            return EnumerationResult(products, failures, total, processed, True)
        processed += 1
        combination_id = f"COMB-{combination_number:07d}"
        try:
            raw_sets = template.reaction.RunReactants(tuple(block.mol for block in blocks_tuple))
        except Exception as exc:
            failures.append(_failure_for_combination(template, combination_id, blocks_tuple, f"reaction error: {exc}"))
            if progress:
                progress(processed, total, len(products))
            continue

        unique_products: dict[str, Chem.Mol] = {}
        rejected = 0
        for product_set in raw_sets:
            if len(product_set) != 1:
                rejected += 1
                continue
            product = Chem.Mol(product_set[0])
            try:
                Chem.SanitizeMol(product)
                smiles = _mol_to_smiles(product)
            except Exception:
                rejected += 1
                continue
            unique_products.setdefault(smiles, product)

        if not unique_products:
            reason = "no product generated"
            if raw_sets and rejected:
                reason += f" ({rejected} invalid or multi-fragment product set(s))"
            failures.append(_failure_for_combination(template, combination_id, blocks_tuple, reason))
        else:
            result_count = len(unique_products)
            for ordinal, (smiles, product) in enumerate(sorted(unique_products.items()), 1):
                product_id = f"{product_prefix}{next_product_number:06d}"
                next_product_number += 1
                first_product_id = first_product_by_smiles.get(smiles, "")
                row: dict[str, Any] = {
                    "Product_ID": product_id,
                    "Product_SMILES": smiles,
                    "Reaction_ID": template.id,
                    "Reaction_Type": template.label,
                    "Combination_ID": combination_id,
                    "Product_Ordinal": ordinal,
                    "Products_From_Combination": result_count,
                    "Ambiguous": "Yes" if result_count > 1 else "No",
                    "Duplicate_Product": "Yes" if first_product_id else "No",
                    "First_Product_ID": first_product_id,
                    "Molecular_Formula": rdMolDescriptors.CalcMolFormula(product),
                    "Molecular_Weight": round(Descriptors.MolWt(product), 4),
                    "Exact_Mass": round(Descriptors.ExactMolWt(product), 6),
                }
                for role_index, block in enumerate(blocks_tuple, 1):
                    row[f"BB{role_index}_Role"] = template.reactant_labels[role_index - 1]
                    _prefix_properties(row, role_index, block)
                products.append(row)
                if not first_product_id:
                    first_product_by_smiles[smiles] = product_id
        if progress:
            progress(processed, total, len(products))

    return EnumerationResult(products, failures, total, processed, False)


def estimate_combinations(template: ReactionTemplate, role_records: Sequence[Sequence[BuildingBlock]]) -> int:
    if len(role_records) != template.reactant_count:
        return 0
    total = 1
    for index, records in enumerate(role_records):
        compatible, _ = compatible_blocks(template, index, records)
        total *= len(compatible)
    return total


def _ordered_columns(rows: Sequence[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                columns.append(key)
    return columns


def _write_csv(path: Path, rows: Sequence[dict[str, Any]]) -> None:
    columns = _ordered_columns(rows)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        if columns:
            writer.writeheader()
            writer.writerows(rows)


def export_csv(path: str | Path, result: EnumerationResult) -> list[Path]:
    destination = Path(path).with_suffix(".csv")
    _write_csv(destination, result.products)
    outputs = [destination]
    if result.failures:
        failure_path = destination.with_name(destination.stem + "_failures.csv")
        _write_csv(failure_path, result.failures)
        outputs.append(failure_path)
    return outputs


def export_excel(path: str | Path, result: EnumerationResult) -> list[Path]:
    destination = Path(path).with_suffix(".xlsx")
    workbook = Workbook()
    products_sheet = workbook.active
    products_sheet.title = "Products"

    def fill_sheet(sheet: Any, rows: Sequence[dict[str, Any]]) -> None:
        columns = _ordered_columns(rows)
        if columns:
            sheet.append(columns)
            for row in rows:
                sheet.append([row.get(column, "") for column in columns])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

    fill_sheet(products_sheet, result.products)
    failures_sheet = workbook.create_sheet("Failures")
    fill_sheet(failures_sheet, result.failures)
    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Total combinations", result.total_combinations])
    summary.append(["Processed combinations", result.processed_combinations])
    summary.append(["Product rows", len(result.products)])
    summary.append(["Unique products", result.unique_product_count])
    summary.append(["Failures/warnings", len(result.failures)])
    summary.append(["Cancelled", "Yes" if result.cancelled else "No"])
    workbook.save(destination)
    return [destination]


def export_sdf(path: str | Path, result: EnumerationResult) -> list[Path]:
    destination = Path(path).with_suffix(".sdf")
    writer = Chem.SDWriter(str(destination))
    try:
        for row in result.products:
            mol = Chem.MolFromSmiles(_clean(row.get("Product_SMILES")))
            if mol is None:
                continue
            mol.SetProp("_Name", _clean(row.get("Product_ID")))
            for key, value in row.items():
                if key != "Product_SMILES":
                    mol.SetProp(str(key), _clean(value))
            writer.write(mol)
    finally:
        writer.close()
    outputs = [destination]
    if result.failures:
        failure_path = destination.with_name(destination.stem + "_failures.csv")
        _write_csv(failure_path, result.failures)
        outputs.append(failure_path)
    return outputs


def export_results(path: str | Path, result: EnumerationResult) -> list[Path]:
    destination = Path(path)
    suffix = destination.suffix.lower()
    if suffix == ".csv":
        return export_csv(destination, result)
    if suffix == ".xlsx":
        return export_excel(destination, result)
    if suffix == ".sdf":
        return export_sdf(destination, result)
    raise ValueError(f"Unsupported output format: {destination.suffix or '(none)'}")
