from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from rdkit import Chem

from bulk_engine import (
    EnumerationResult,
    enumerate_library,
    export_results,
    inspect_source,
    load_building_blocks,
    load_reaction_templates,
    single_building_block,
)


class BulkEngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.templates = {item.id: item for item in load_reaction_templates()}

    def test_registry_has_all_templates_and_role_counts(self) -> None:
        self.assertEqual(len(self.templates), 43)
        self.assertEqual(self.templates["petasis"].reactant_count, 3)
        self.assertEqual(self.templates["amide_coupling"].reactant_count, 2)
        self.assertEqual(self.templates["nitro_reduction"].reactant_count, 1)

    def test_csv_load_and_amide_lineage(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "acids.csv"
            path.write_text("ID,SMILES,Vendor\nA1,CC(=O)O,Alpha\nA2,CCC(=O)O,Beta\nBAD,not-smiles,Gamma\n", encoding="utf-8")
            inspection = inspect_source(path)
            self.assertEqual(inspection.default_id_field, "ID")
            loaded = load_building_blocks(path, "ID", "SMILES")
            self.assertEqual(len(loaded.records), 2)
            self.assertEqual(len(loaded.issues), 1)

            amine = single_building_block("N1", "CN")
            result = enumerate_library(
                self.templates["amide_coupling"],
                [loaded.records, [amine]],
                prior_failures=loaded.issues,
                product_prefix="DOCK-",
                start_number=10,
            )
            self.assertEqual(result.total_combinations, 2)
            self.assertEqual(len(result.products), 2)
            self.assertEqual(result.products[0]["Product_ID"], "DOCK-000010")
            self.assertEqual({row["BB1_ID"] for row in result.products}, {"A1", "A2"})
            self.assertEqual({row["BB1_Vendor"] for row in result.products}, {"Alpha", "Beta"})
            self.assertEqual({row["BB2_ID"] for row in result.products}, {"N1"})

    def test_incompatible_blocks_are_reported(self) -> None:
        acid = single_building_block("A1", "CC(=O)O")
        not_amine = single_building_block("X1", "CC")
        result = enumerate_library(self.templates["amide_coupling"], [[acid], [not_amine]])
        self.assertEqual(result.total_combinations, 0)
        self.assertEqual(result.products, [])
        self.assertTrue(any(row["Failure_Type"] == "Compatibility" for row in result.failures))

    def test_sdf_input_uses_structure_and_properties(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "blocks.sdf"
            writer = Chem.SDWriter(str(path))
            mol = Chem.MolFromSmiles("CN")
            mol.SetProp("_Name", "AM-1")
            mol.SetProp("Vendor", "Supplier A")
            writer.write(mol)
            writer.close()

            inspection = inspect_source(path)
            self.assertIn("_Name", inspection.fields)
            self.assertIn("Vendor", inspection.fields)
            loaded = load_building_blocks(path, "_Name")
            self.assertEqual(loaded.records[0].identifier, "AM-1")
            self.assertEqual(loaded.records[0].properties["Vendor"], "Supplier A")

    def test_all_export_formats(self) -> None:
        row = {
            "Product_ID": "P-1",
            "Product_SMILES": "CC",
            "BB1_ID": "A",
            "BB2_ID": "B",
        }
        result = EnumerationResult([row], [{"Failure_Type": "Input", "Reason": "example"}], 1, 1)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_outputs = export_results(root / "out.csv", result)
            self.assertEqual(len(csv_outputs), 2)
            with csv_outputs[0].open(encoding="utf-8-sig") as handle:
                self.assertEqual(next(csv.DictReader(handle))["BB2_ID"], "B")

            xlsx = export_results(root / "out.xlsx", result)[0]
            workbook = load_workbook(xlsx, read_only=True)
            self.assertEqual(workbook.sheetnames, ["Products", "Failures", "Summary"])
            workbook.close()

            sdf_outputs = export_results(root / "out.sdf", result)
            supplier = Chem.SDMolSupplier(str(sdf_outputs[0]))
            self.assertEqual(supplier[0].GetProp("BB1_ID"), "A")
            del supplier


if __name__ == "__main__":
    unittest.main()
