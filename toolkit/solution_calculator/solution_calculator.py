from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Dict, Iterable, List, Optional

from openpyxl import load_workbook
from PIL import ImageTk

from solution_calculations import (
    check_concentration,
    formula_molecular_weight,
    format_number,
    positive_number,
    prepare_stock,
    required_amount,
)

try:
    from rdkit import Chem
    from rdkit.Chem import Descriptors, Draw, rdMolDescriptors
except ImportError:
    Chem = None
    Descriptors = None
    Draw = None
    rdMolDescriptors = None


PAYLOAD_SCHEMA = "xinpilot_solution_calculator.v1"
MODE_PREPARE = "Prepare stock from available amount"
MODE_CHECK = "Check concentration from amount and volume"
MODE_REQUIRED = "Calculate amount required for a target volume"
IMPORT_SMILES = "SMILES"
IMPORT_FORMULA = "Molecular formula"
IMPORT_MW = "Molecular weight"
SDF_MOLECULE_COLUMN = "SDF molecule"


def _amount_unit(value: str) -> str:
    return "umol" if value in {"umol", "µmol", "μmol"} else value


def _concentration_unit(value: str) -> str:
    return "uM" if value in {"uM", "µM", "μM"} else value


def _volume_unit(value: str) -> str:
    return "uL" if value in {"uL", "µL", "μL"} else value


def load_xinpilot_payload(path: object) -> Dict[str, object]:
    payload_path = Path(path)
    payload: Dict[str, object] = {}
    delete_after_load = False
    try:
        payload = json.loads(payload_path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict) or payload.get("schema") not in (PAYLOAD_SCHEMA, "rxnpilot_solution_calculator.v1", "syntrack_solution_calculator.v1"):
            raise ValueError("This is not a supported XinPilot Solution Calculator transfer file.")
        products = payload.get("products", [])
        if not isinstance(products, list):
            raise ValueError("The transfer file has an invalid products list.")
        delete_after_load = bool(payload.get("delete_after_load", False))
        return payload
    finally:
        if delete_after_load:
            try:
                payload_path.unlink()
            except OSError:
                pass


def product_to_batch_row(product: Dict[str, object]) -> Dict[str, str]:
    reaction_well = str(product.get("reaction_well", "") or "").strip()
    product_id = str(product.get("product_id", "") or "").strip()
    product_name = str(product.get("product_name", "") or product_id or reaction_well).strip()
    return {
        "reaction_well": reaction_well,
        "lcms_well": str(product.get("lcms_well", "") or "").strip(),
        "reaction_name": str(product.get("reaction_name", "") or "").strip(),
        "product_id": product_id,
        "product_name": product_name,
        "smiles": str(product.get("product_smiles", "") or "").strip(),
        "formula": str(product.get("product_formula", "") or "").strip(),
        "exact_mass": str(product.get("product_exact_mass", "") or "").strip(),
        "mw": str(product.get("product_mw", "") or "").strip(),
        "lcms_reference_smiles": str(product.get("lcms_reference_smiles", "") or "").strip(),
        "uv_purity_reference": str(product.get("purity", "") or "").strip(),
        "uv_purity_source": str(product.get("purity_source", "") or "").strip(),
        "amount": "",
        "amount_unit": "mg",
        "mass_purity": "100",
        "target_mm": "10",
        "solvent": "DMSO",
        "rounding_ul": "0",
        "exact_volume_ul": "",
        "final_volume_ul": "",
        "achieved_mm": "",
        "status": "Enter recovered amount",
        "identity_source": "XinPilot Product Details SMILES",
        "identity_value": str(product.get("product_smiles", "") or "").strip(),
        "identity_error": "",
    }


def read_batch_import_file(path: object, sheet_name: Optional[str] = None) -> tuple[List[Dict[str, str]], List[str]]:
    source_path = Path(path)
    suffix = source_path.suffix.lower()
    if suffix == ".csv":
        with source_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            columns = [str(value or "").strip() for value in (reader.fieldnames or []) if str(value or "").strip()]
            if not columns:
                raise ValueError("The CSV file does not contain a header row.")
            rows = [
                {column: str(row.get(column, "") or "").strip() for column in columns}
                for row in reader
            ]
        return rows, columns
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        try:
            selected_sheet = sheet_name or workbook.sheetnames[0]
            if selected_sheet not in workbook.sheetnames:
                raise ValueError(f"Worksheet not found: {selected_sheet}")
            worksheet = workbook[selected_sheet]
            iterator = worksheet.iter_rows(values_only=True)
            try:
                header_values = next(iterator)
            except StopIteration as exc:
                raise ValueError("The selected worksheet is empty.") from exc
            columns = []
            column_indices = []
            for index, value in enumerate(header_values):
                heading = str(value or "").strip()
                if heading and heading not in columns:
                    columns.append(heading)
                    column_indices.append(index)
            if not columns:
                raise ValueError("The selected worksheet does not contain a header row.")
            rows = []
            for values in iterator:
                row = {
                    column: str(values[index] if index < len(values) and values[index] is not None else "").strip()
                    for column, index in zip(columns, column_indices)
                }
                if any(row.values()):
                    rows.append(row)
            return rows, columns
        finally:
            workbook.close()
    if suffix in {".sdf", ".sd"}:
        if Chem is None:
            raise ValueError("RDKit is required to import SDF files.")
        supplier = Chem.SDMolSupplier(str(source_path), removeHs=False)
        rows: List[Dict[str, str]] = []
        columns = ["_Name", SDF_MOLECULE_COLUMN]
        for molecule in supplier:
            if molecule is None:
                rows.append({"_Name": "", SDF_MOLECULE_COLUMN: ""})
                continue
            row = {
                "_Name": str(molecule.GetProp("_Name") if molecule.HasProp("_Name") else "").strip(),
                SDF_MOLECULE_COLUMN: Chem.MolToSmiles(molecule, canonical=False, isomericSmiles=True),
            }
            for name in molecule.GetPropNames():
                key = str(name)
                row[key] = str(molecule.GetProp(name) or "").strip()
                if key not in columns:
                    columns.append(key)
            rows.append(row)
        if not rows:
            raise ValueError("The SDF file does not contain any readable records.")
        return rows, columns
    raise ValueError("Choose an SDF, Excel (.xlsx), or CSV file.")


def build_import_batch_rows(
    records: Iterable[Dict[str, object]],
    *,
    id_column: str,
    source_type: str,
    source_column: str,
    target_mm: object,
    solvent: object,
    amount_unit: str,
) -> List[Dict[str, str]]:
    if source_type not in {IMPORT_SMILES, IMPORT_FORMULA, IMPORT_MW}:
        raise ValueError(f"Unsupported calculation source: {source_type}")
    imported: List[Dict[str, str]] = []
    for index, record in enumerate(records, start=1):
        identifier = str(record.get(id_column, "") or "").strip()
        source_value = str(record.get(source_column, "") or "").strip()
        errors = []
        if not identifier:
            errors.append(f"ID is blank in row {index}")
            identifier = f"Row {index}"
        product = {"product_id": identifier, "product_name": identifier}
        try:
            if source_type == IMPORT_SMILES:
                if Chem is None:
                    raise ValueError("RDKit is required to use SMILES")
                molecule = Chem.MolFromSmiles(source_value) if source_value else None
                if molecule is None:
                    raise ValueError("invalid or blank SMILES")
                product.update({
                    "product_smiles": source_value,
                    "product_formula": rdMolDescriptors.CalcMolFormula(molecule),
                    "product_exact_mass": f"{rdMolDescriptors.CalcExactMolWt(molecule):.6g}",
                    "product_mw": f"{Descriptors.MolWt(molecule):.6g}",
                })
            elif source_type == IMPORT_FORMULA:
                mw = formula_molecular_weight(source_value)
                product.update({"product_formula": source_value, "product_mw": f"{mw:.6g}"})
            else:
                mw = positive_number(source_value, "Molecular weight")
                product.update({"product_mw": f"{mw:.6g}"})
        except ValueError as exc:
            errors.append(str(exc))

        row = product_to_batch_row(product)
        row.update({
            "target_mm": str(target_mm or "").strip(),
            "solvent": str(solvent or "").strip(),
            "amount_unit": amount_unit,
            "identity_source": source_type,
            "identity_value": source_value,
            "identity_error": "; ".join(errors),
            "status": "; ".join(errors) if errors else "Enter recovered amount",
        })
        imported.append(row)
    return imported


class SolutionCalculatorApp(tk.Tk):
    def __init__(self, payload: Optional[Dict[str, object]] = None):
        super().__init__()
        self.title("Solution & Stock Calculator")
        self.geometry("1280x800")
        self.minsize(1000, 680)
        self.batch_rows: List[Dict[str, str]] = []
        self.import_batch_rows: List[Dict[str, str]] = []
        self._batch_preview_window: Optional[tk.Toplevel] = None
        self._batch_preview_image = None
        self._build_ui()
        self._init_selection_editors()
        products = (payload or {}).get("products", [])
        if isinstance(products, list) and products:
            self.add_products(item for item in products if isinstance(item, dict))
            self.notebook.select(self.batch_tab)

    def _build_ui(self) -> None:
        outer = ttk.Frame(self, padding=12)
        outer.pack(fill="both", expand=True)
        ttk.Label(outer, text="Solution & Stock Calculator", font=("Segoe UI", 17, "bold")).pack(anchor="w")
        ttk.Label(
            outer,
            text=(
                "Calculate preparation volumes, resulting concentrations, or required mass. "
                "Results use final solution volume: dissolve the compound, then make up to the stated final volume."
            ),
            wraplength=1120,
            justify="left",
        ).pack(anchor="w", pady=(3, 10))

        self.notebook = ttk.Notebook(outer)
        self.notebook.pack(fill="both", expand=True)
        self.single_tab = ttk.Frame(self.notebook, padding=14)
        self.batch_tab = ttk.Frame(self.notebook, padding=10)
        self.import_tab = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.single_tab, text="Single Compound")
        self.notebook.add(self.batch_tab, text="Batch from XinPilot")
        self.notebook.add(self.import_tab, text="Import Batch")
        self._build_single_tab()
        self._build_batch_tab()
        self._build_import_tab()

    def _build_single_tab(self) -> None:
        self.single_tab.columnconfigure(0, weight=1)
        self.single_tab.columnconfigure(1, weight=1)
        self.single_tab.rowconfigure(2, weight=1)

        mode_box = ttk.LabelFrame(self.single_tab, text="Calculation", padding=10)
        mode_box.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 10))
        mode_box.columnconfigure(1, weight=1)
        self.mode_var = tk.StringVar(value=MODE_PREPARE)
        ttk.Label(mode_box, text="Mode").grid(row=0, column=0, sticky="w", padx=(0, 8))
        mode = ttk.Combobox(
            mode_box,
            textvariable=self.mode_var,
            values=(MODE_PREPARE, MODE_CHECK, MODE_REQUIRED),
            state="readonly",
            width=50,
        )
        mode.grid(row=0, column=1, sticky="w")
        mode.bind("<<ComboboxSelected>>", self._update_single_mode_hint)
        self.mode_hint_var = tk.StringVar()
        ttk.Label(mode_box, textvariable=self.mode_hint_var, wraplength=920, justify="left").grid(
            row=1, column=0, columnspan=3, sticky="w", pady=(8, 0)
        )

        identity = ttk.LabelFrame(self.single_tab, text="Compound", padding=10)
        identity.grid(row=1, column=0, sticky="nsew", padx=(0, 5), pady=(0, 10))
        inputs = ttk.LabelFrame(self.single_tab, text="Amounts and conditions", padding=10)
        inputs.grid(row=1, column=1, sticky="nsew", padx=(5, 0), pady=(0, 10))
        for frame in (identity, inputs):
            frame.columnconfigure(1, weight=1)

        self.single_name_var = tk.StringVar()
        self.single_smiles_var = tk.StringVar()
        self.single_formula_var = tk.StringVar()
        self.single_mw_var = tk.StringVar()
        self.single_exact_mass_var = tk.StringVar()
        self.single_solvent_var = tk.StringVar(value="DMSO")
        self._entry_row(identity, 0, "Compound name", self.single_name_var)
        self._entry_row(identity, 1, "SMILES", self.single_smiles_var)
        ttk.Button(identity, text="Calculate identity from SMILES", command=self._identity_from_smiles).grid(
            row=2, column=1, sticky="w", pady=(2, 7)
        )
        self._entry_row(identity, 3, "Formula", self.single_formula_var)
        self._entry_row(identity, 4, "Molecular weight (g/mol)", self.single_mw_var)
        self._entry_row(identity, 5, "Exact mass", self.single_exact_mass_var)
        self._entry_row(identity, 6, "Solvent", self.single_solvent_var)

        self.single_amount_var = tk.StringVar()
        self.single_amount_unit_var = tk.StringVar(value="mg")
        self.single_purity_var = tk.StringVar(value="100")
        self.single_concentration_var = tk.StringVar(value="10")
        self.single_concentration_unit_var = tk.StringVar(value="mM")
        self.single_volume_var = tk.StringVar(value="1")
        self.single_volume_unit_var = tk.StringVar(value="mL")
        self.single_rounding_var = tk.StringVar(value="0")
        self.single_amount_widgets = self._value_unit_row(inputs, 0, "Available amount", self.single_amount_var, self.single_amount_unit_var, ("mg", "µmol"))
        self._entry_row(inputs, 1, "Mass/assay purity (%)", self.single_purity_var)
        self.single_concentration_widgets = self._value_unit_row(
            inputs, 2, "Target concentration", self.single_concentration_var, self.single_concentration_unit_var, ("µM", "mM", "M")
        )
        self.single_volume_widgets = self._value_unit_row(inputs, 3, "Final solution volume", self.single_volume_var, self.single_volume_unit_var, ("µL", "mL"))
        self.single_rounding_entry = self._entry_row(inputs, 4, "Round prepared volume to (µL)", self.single_rounding_var)
        ttk.Label(
            inputs,
            text="Use mass/assay purity only when it is known. UV254 area purity is not the same measurement.",
            wraplength=470,
            justify="left",
        ).grid(row=5, column=0, columnspan=3, sticky="w", pady=(8, 0))

        results = ttk.LabelFrame(self.single_tab, text="Result and preparation note", padding=10)
        results.grid(row=2, column=0, columnspan=2, sticky="nsew")
        results.columnconfigure(0, weight=1)
        results.rowconfigure(1, weight=1)
        ttk.Button(results, text="Calculate", command=self._calculate_single).grid(row=0, column=0, sticky="w", pady=(0, 8))
        self.single_copy_button = ttk.Button(results, text="Copy result", command=self._copy_single_result, state="disabled")
        self.single_copy_button.grid(row=0, column=1, sticky="e", pady=(0, 8))
        self.single_result = tk.Text(results, height=9, wrap="word", font=("Segoe UI", 10))
        self.single_result.grid(row=1, column=0, columnspan=2, sticky="nsew")
        self._set_text(self.single_result, "Enter the values above, then select Calculate.")
        self._single_result_valid = False
        for variable in (
            self.mode_var, self.single_name_var, self.single_smiles_var, self.single_formula_var,
            self.single_mw_var, self.single_exact_mass_var, self.single_solvent_var,
            self.single_amount_var, self.single_amount_unit_var, self.single_purity_var,
            self.single_concentration_var, self.single_concentration_unit_var,
            self.single_volume_var, self.single_volume_unit_var, self.single_rounding_var,
        ):
            variable.trace_add("write", self._invalidate_single_result)
        self.mode_var.trace_add("write", lambda *_args: self._update_single_mode_hint())
        self._update_single_mode_hint()

    def _invalidate_single_result(self, *_args):
        self._single_result_valid = False
        self.single_copy_button.configure(state="disabled")
        self._set_text(self.single_result, "Inputs changed. Select Calculate to obtain a current result.")

    @staticmethod
    def _entry_row(parent: ttk.Frame, row: int, label: str, variable: tk.StringVar) -> ttk.Entry:
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(0, 8), pady=3)
        entry = ttk.Entry(parent, textvariable=variable)
        entry.grid(row=row, column=1, columnspan=2, sticky="ew", pady=3)
        return entry

    @staticmethod
    def _value_unit_row(
        parent: ttk.Frame,
        row: int,
        label: str,
        value_var: tk.StringVar,
        unit_var: tk.StringVar,
        units: Iterable[str],
    ) -> tuple[ttk.Entry, ttk.Combobox]:
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(0, 8), pady=3)
        entry = ttk.Entry(parent, textvariable=value_var)
        entry.grid(row=row, column=1, sticky="ew", pady=3)
        combo = ttk.Combobox(parent, textvariable=unit_var, values=tuple(units), state="readonly", width=8)
        combo.grid(row=row, column=2, sticky="w", padx=(6, 0), pady=3)
        return entry, combo

    def _update_single_mode_hint(self, _event=None) -> None:
        hints = {
            MODE_PREPARE: "Uses available mg or µmol and target concentration to calculate the final stock-solution volume.",
            MODE_CHECK: "Uses available mg or µmol and an actual final solution volume to calculate concentration.",
            MODE_REQUIRED: "Uses target concentration and desired final volume to calculate the required µmol and mg.",
        }
        mode = self.mode_var.get()
        self.mode_hint_var.set(hints[mode])
        for widgets, enabled in (
            (self.single_amount_widgets, mode != MODE_REQUIRED),
            (self.single_concentration_widgets, mode != MODE_CHECK),
            (self.single_volume_widgets, mode != MODE_PREPARE),
        ):
            widgets[0].configure(state="normal" if enabled else "disabled")
            widgets[1].configure(state="readonly" if enabled else "disabled")
        self.single_rounding_entry.configure(state="normal" if mode == MODE_PREPARE else "disabled")

    def _identity_from_smiles(self) -> None:
        if Chem is None:
            messagebox.showerror("RDKit unavailable", "RDKit is required to calculate identity from SMILES.", parent=self)
            return
        molecule = Chem.MolFromSmiles(self.single_smiles_var.get().strip())
        if molecule is None:
            messagebox.showerror("Invalid SMILES", "The entered SMILES could not be parsed.", parent=self)
            return
        self.single_smiles_var.set(Chem.MolToSmiles(molecule, canonical=True, isomericSmiles=True))
        self.single_formula_var.set(rdMolDescriptors.CalcMolFormula(molecule))
        self.single_mw_var.set(f"{Descriptors.MolWt(molecule):.4f}")
        self.single_exact_mass_var.set(f"{rdMolDescriptors.CalcExactMolWt(molecule):.4f}")

    def _calculate_single(self) -> None:
        self._invalidate_single_result()
        try:
            mode = self.mode_var.get()
            name = self.single_name_var.get().strip() or "Compound"
            solvent = self.single_solvent_var.get().strip() or "selected solvent"
            if mode == MODE_PREPARE:
                result = prepare_stock(
                    amount=self.single_amount_var.get(),
                    amount_unit=_amount_unit(self.single_amount_unit_var.get()),
                    target_concentration=self.single_concentration_var.get(),
                    concentration_unit=_concentration_unit(self.single_concentration_unit_var.get()),
                    molecular_weight=self.single_mw_var.get(),
                    mass_purity_percent=self.single_purity_var.get(),
                    rounding_increment_ul=self.single_rounding_var.get(),
                )
                lines = [
                    f"{name}: {format_number(result.active_umol)} µmol active compound is available.",
                    f"Exact final volume: {format_number(result.exact_volume_ul)} µL.",
                ]
                if abs(result.final_volume_ul - result.exact_volume_ul) > 1e-9:
                    lines.append(
                        f"Rounded final volume: {format_number(result.final_volume_ul)} µL; achieved concentration: "
                        f"{format_number(result.achieved_concentration_mm)} mM."
                    )
                lines.append(
                    f"Preparation: dissolve in {solvent}, then make up to a final solution volume of "
                    f"{format_number(result.final_volume_ul)} µL. Do not interpret this as adding that full volume to a solid."
                )
            elif mode == MODE_CHECK:
                result = check_concentration(
                    amount=self.single_amount_var.get(),
                    amount_unit=_amount_unit(self.single_amount_unit_var.get()),
                    final_volume=self.single_volume_var.get(),
                    volume_unit=_volume_unit(self.single_volume_unit_var.get()),
                    molecular_weight=self.single_mw_var.get(),
                    mass_purity_percent=self.single_purity_var.get(),
                )
                lines = [
                    f"{name}: {format_number(result.active_umol)} µmol active compound in the stated final volume.",
                    f"Calculated concentration: {format_number(result.concentration_mm)} mM.",
                ]
            else:
                result = required_amount(
                    target_concentration=self.single_concentration_var.get(),
                    concentration_unit=_concentration_unit(self.single_concentration_unit_var.get()),
                    final_volume=self.single_volume_var.get(),
                    volume_unit=_volume_unit(self.single_volume_unit_var.get()),
                    molecular_weight=self.single_mw_var.get(),
                    mass_purity_percent=self.single_purity_var.get(),
                )
                lines = [
                    f"{name}: {format_number(result.active_umol)} µmol active compound is required.",
                    f"At the stated mass/assay purity, weigh {format_number(result.required_mg)} mg "
                    f"({format_number(result.weighed_umol)} µmol as weighed).",
                    f"Dissolve in {solvent}, then make up to the requested final solution volume.",
                ]
            self._set_text(self.single_result, "\n\n".join(lines))
            self._single_result_valid = True
            self.single_copy_button.configure(state="normal")
        except ValueError as exc:
            self._set_text(self.single_result, f"Cannot calculate: {exc}")
            messagebox.showerror("Cannot calculate", str(exc), parent=self)

    def _copy_single_result(self) -> None:
        if not self._single_result_valid:
            return
        text = self.single_result.get("1.0", "end-1c").strip()
        if text:
            self.clipboard_clear()
            self.clipboard_append(text)

    def _build_batch_tab(self) -> None:
        self.batch_tab.columnconfigure(0, weight=1)
        self.batch_tab.rowconfigure(1, weight=1)

        common = ttk.LabelFrame(self.batch_tab, text="Common settings", padding=8)
        common.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        self.batch_common_target_var = tk.StringVar(value="10")
        self.batch_common_solvent_var = tk.StringVar(value="DMSO")
        self.batch_common_rounding_var = tk.StringVar(value="0")
        ttk.Label(common, text="Target (mM)").pack(side="left")
        ttk.Entry(common, textvariable=self.batch_common_target_var, width=9).pack(side="left", padx=(5, 12))
        ttk.Label(common, text="Solvent").pack(side="left")
        ttk.Entry(common, textvariable=self.batch_common_solvent_var, width=14).pack(side="left", padx=(5, 12))
        ttk.Label(common, text="Round volume to (µL)").pack(side="left")
        ttk.Entry(common, textvariable=self.batch_common_rounding_var, width=9).pack(side="left", padx=(5, 12))
        ttk.Button(common, text="Apply to all rows", command=self._apply_common_settings).pack(side="left")

        table_box = ttk.Frame(self.batch_tab)
        table_box.grid(row=1, column=0, sticky="nsew")
        table_box.columnconfigure(0, weight=1)
        table_box.rowconfigure(0, weight=1)
        columns = (
            "rxn", "product", "formula", "mw", "amount", "unit", "purity", "target", "solvent", "volume", "achieved", "status"
        )
        self.batch_tree = ttk.Treeview(table_box, columns=columns, show="headings", selectmode="extended", height=13)
        for key, label, width in (
            ("rxn", "RXN well", 75), ("product", "Product", 150), ("formula", "Formula", 90),
            ("mw", "MW", 72), ("amount", "Recovered", 78), ("unit", "Unit", 50),
            ("purity", "Mass purity %", 85), ("target", "Target mM", 76), ("solvent", "Solvent", 80),
            ("volume", "Final µL", 80), ("achieved", "Achieved mM", 88), ("status", "Status", 180),
        ):
            self.batch_tree.heading(key, text=label)
            self.batch_tree.column(key, width=width, minwidth=45, anchor="w")
        ybar = ttk.Scrollbar(table_box, orient="vertical", command=self.batch_tree.yview)
        xbar = ttk.Scrollbar(table_box, orient="horizontal", command=self.batch_tree.xview)
        self.batch_tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        self.batch_tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        self.batch_tree.bind("<<TreeviewSelect>>", self._load_selected_batch_row)

        editor = ttk.LabelFrame(self.batch_tab, text="Selected row(s) - only changed fields are applied", padding=8)
        editor.grid(row=2, column=0, sticky="ew", pady=(8, 0))
        for column in range(12):
            editor.columnconfigure(column, weight=1 if column in {1, 3, 5, 7, 9} else 0)
        self.batch_amount_var = tk.StringVar()
        self.batch_amount_unit_var = tk.StringVar(value="mg")
        self.batch_mass_purity_var = tk.StringVar(value="100")
        self.batch_target_var = tk.StringVar(value="10")
        self.batch_solvent_var = tk.StringVar(value="DMSO")
        self.batch_rounding_var = tk.StringVar(value="0")
        pairs = (
            ("Recovered amount", self.batch_amount_var), ("Mass purity %", self.batch_mass_purity_var),
            ("Target mM", self.batch_target_var), ("Solvent", self.batch_solvent_var), ("Round µL", self.batch_rounding_var),
        )
        column = 0
        for label, variable in pairs:
            ttk.Label(editor, text=label).grid(row=0, column=column, sticky="w", padx=(0, 4))
            entry = ttk.Entry(editor, textvariable=variable, width=11)
            entry.grid(row=0, column=column + 1, sticky="ew", padx=(0, 10))
            if label == "Recovered amount":
                self.batch_amount_entry = entry
            column += 2
        ttk.Combobox(editor, textvariable=self.batch_amount_unit_var, values=("mg", "µmol"), state="readonly", width=7).grid(
            row=0, column=10, sticky="w"
        )
        self.batch_selected_actions = ttk.Frame(editor)
        self.batch_selected_actions.grid(row=1, column=1, columnspan=5, sticky="w", pady=(7, 0))
        self.batch_update_selected_button = ttk.Button(
            self.batch_selected_actions, text="Update selected", command=self._save_selected_batch_row
        )
        self.batch_update_selected_button.pack(side="left")
        self.batch_calculate_all_button = ttk.Button(
            self.batch_selected_actions, text="Calculate all", command=self._calculate_all_rows
        )
        self.batch_calculate_all_button.pack(side="left", padx=(6, 0))

        actions = ttk.Frame(self.batch_tab)
        actions.grid(row=3, column=0, sticky="ew", pady=(8, 0))
        self.batch_preview_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            actions,
            text="Show structure preview",
            variable=self.batch_preview_var,
            command=self._toggle_batch_structure_preview,
        ).pack(side="left", padx=(0, 12))
        ttk.Button(actions, text="Add blank row", command=self._add_blank_batch_row).pack(side="left")
        ttk.Button(actions, text="Remove selected", command=self._remove_selected_rows).pack(side="left", padx=(6, 0))
        ttk.Button(actions, text="Use selected in Single Compound", command=self._send_selected_to_single).pack(side="left", padx=(6, 0))
        ttk.Button(actions, text="Export CSV", command=self._export_batch_csv).pack(side="right")
        ttk.Button(actions, text="Copy preparation notes", command=self._copy_batch_notes).pack(side="right", padx=(0, 6))

    def _build_import_tab(self) -> None:
        self.import_tab.columnconfigure(0, weight=1)
        self.import_tab.rowconfigure(2, weight=1)

        source_box = ttk.LabelFrame(self.import_tab, text="Import source", padding=8)
        source_box.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        self.import_file_var = tk.StringVar(value="No file selected")
        ttk.Button(source_box, text="Choose SDF, Excel, or CSV…", command=self._choose_import_batch_file).pack(side="left")
        ttk.Label(source_box, textvariable=self.import_file_var).pack(side="left", padx=(10, 0))

        common = ttk.LabelFrame(self.import_tab, text="Batch settings", padding=8)
        common.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        self.import_common_target_var = tk.StringVar(value="10")
        self.import_common_solvent_var = tk.StringVar(value="DMSO")
        self.import_common_amount_unit_var = tk.StringVar(value="mg")
        self.import_common_rounding_var = tk.StringVar(value="0")
        ttk.Label(common, text="Target (mM)").pack(side="left")
        ttk.Entry(common, textvariable=self.import_common_target_var, width=9).pack(side="left", padx=(5, 12))
        ttk.Label(common, text="Solvent").pack(side="left")
        ttk.Entry(common, textvariable=self.import_common_solvent_var, width=14).pack(side="left", padx=(5, 12))
        ttk.Label(common, text="Recovered amount unit").pack(side="left")
        ttk.Combobox(
            common,
            textvariable=self.import_common_amount_unit_var,
            values=("mg", "µmol"),
            state="readonly",
            width=7,
        ).pack(side="left", padx=(5, 12))
        ttk.Label(common, text="Round volume to (µL)").pack(side="left")
        ttk.Entry(common, textvariable=self.import_common_rounding_var, width=9).pack(side="left", padx=(5, 12))
        ttk.Button(common, text="Apply to all rows", command=self._apply_import_common_settings).pack(side="left")

        table_box = ttk.Frame(self.import_tab)
        table_box.grid(row=2, column=0, sticky="nsew")
        table_box.columnconfigure(0, weight=1)
        table_box.rowconfigure(0, weight=1)
        columns = ("id", "source", "formula", "mw", "amount", "unit", "purity", "target", "solvent", "volume", "achieved", "status")
        self.import_tree = ttk.Treeview(table_box, columns=columns, show="headings", selectmode="extended", height=13)
        for key, label, width in (
            ("id", "ID", 150), ("source", "Calculation source", 120), ("formula", "Formula", 90),
            ("mw", "MW", 72), ("amount", "Recovered", 78), ("unit", "Unit", 50),
            ("purity", "Mass purity %", 85), ("target", "Target mM", 76), ("solvent", "Solvent", 80),
            ("volume", "Final µL", 80), ("achieved", "Achieved mM", 88), ("status", "Status", 190),
        ):
            self.import_tree.heading(key, text=label)
            self.import_tree.column(key, width=width, minwidth=45, anchor="w")
        ybar = ttk.Scrollbar(table_box, orient="vertical", command=self.import_tree.yview)
        xbar = ttk.Scrollbar(table_box, orient="horizontal", command=self.import_tree.xview)
        self.import_tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        self.import_tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        self.import_tree.bind("<<TreeviewSelect>>", self._load_selected_import_row)

        editor = ttk.LabelFrame(self.import_tab, text="Selected row(s) - only changed fields are applied", padding=8)
        editor.grid(row=3, column=0, sticky="ew", pady=(8, 0))
        for column in range(12):
            editor.columnconfigure(column, weight=1 if column in {1, 3, 5, 7, 9} else 0)
        self.import_amount_var = tk.StringVar()
        self.import_amount_unit_var = tk.StringVar(value="mg")
        self.import_mass_purity_var = tk.StringVar(value="100")
        self.import_target_var = tk.StringVar(value="10")
        self.import_solvent_var = tk.StringVar(value="DMSO")
        self.import_rounding_var = tk.StringVar(value="0")
        pairs = (
            ("Recovered amount", self.import_amount_var), ("Mass purity %", self.import_mass_purity_var),
            ("Target mM", self.import_target_var), ("Solvent", self.import_solvent_var), ("Round µL", self.import_rounding_var),
        )
        column = 0
        for label, variable in pairs:
            ttk.Label(editor, text=label).grid(row=0, column=column, sticky="w", padx=(0, 4))
            entry = ttk.Entry(editor, textvariable=variable, width=11)
            entry.grid(row=0, column=column + 1, sticky="ew", padx=(0, 10))
            if label == "Recovered amount":
                self.import_amount_entry = entry
            column += 2
        ttk.Combobox(
            editor,
            textvariable=self.import_amount_unit_var,
            values=("mg", "µmol"),
            state="readonly",
            width=7,
        ).grid(row=0, column=10, sticky="w")
        self.import_selected_actions = ttk.Frame(editor)
        self.import_selected_actions.grid(row=1, column=1, columnspan=5, sticky="w", pady=(7, 0))
        self.import_update_selected_button = ttk.Button(
            self.import_selected_actions, text="Update selected", command=self._save_selected_import_row
        )
        self.import_update_selected_button.pack(side="left")
        self.import_calculate_all_button = ttk.Button(
            self.import_selected_actions, text="Calculate all", command=self._calculate_all_import_rows
        )
        self.import_calculate_all_button.pack(side="left", padx=(6, 0))

        actions = ttk.Frame(self.import_tab)
        actions.grid(row=4, column=0, sticky="ew", pady=(8, 0))
        self.import_preview_var = tk.BooleanVar(value=False)
        self.import_preview_check = ttk.Checkbutton(
            actions,
            text="Show structure preview",
            variable=self.import_preview_var,
            command=self._toggle_import_structure_preview,
            state="disabled",
        )
        self.import_preview_check.pack(side="left", padx=(0, 12))
        ttk.Button(actions, text="Remove selected", command=self._remove_selected_import_rows).pack(side="left")
        ttk.Button(actions, text="Use selected in Single Compound", command=self._send_import_selected_to_single).pack(side="left", padx=(6, 0))
        ttk.Button(actions, text="Export CSV", command=self._export_import_batch_csv).pack(side="right")

        self._import_preview_window: Optional[tk.Toplevel] = None
        self._import_preview_image = None

    def _choose_excel_sheet(self, path: Path) -> Optional[str]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = list(workbook.sheetnames)
        finally:
            workbook.close()
        if len(sheets) == 1:
            return sheets[0]
        window = tk.Toplevel(self)
        window.title("Select Excel worksheet")
        window.transient(self)
        window.resizable(False, False)
        window.grab_set()
        body = ttk.Frame(window, padding=14)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="Worksheet").grid(row=0, column=0, sticky="w", padx=(0, 8))
        value = tk.StringVar(value=sheets[0])
        ttk.Combobox(body, textvariable=value, values=sheets, state="readonly", width=32).grid(row=0, column=1, sticky="ew")
        result: Dict[str, str] = {}

        def accept():
            result["sheet"] = value.get()
            window.destroy()

        controls = ttk.Frame(body)
        controls.grid(row=1, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(controls, text="Continue", command=accept).pack(side="left")
        ttk.Button(controls, text="Cancel", command=window.destroy).pack(side="left", padx=(6, 0))
        window.bind("<Return>", lambda _event: accept())
        window.bind("<Escape>", lambda _event: window.destroy())
        self.wait_window(window)
        return result.get("sheet")

    def _choose_import_mapping(self, columns: List[str]) -> Optional[Dict[str, str]]:
        if not columns:
            return None
        lower = {column.casefold(): column for column in columns}
        id_default = next((lower[key] for key in ("id", "identifier", "name", "_name") if key in lower), columns[0])

        def suggested_column(source_type: str) -> str:
            if source_type == IMPORT_SMILES:
                if SDF_MOLECULE_COLUMN in columns:
                    return SDF_MOLECULE_COLUMN
                return next((column for column in columns if "smiles" in column.casefold()), columns[0])
            if source_type == IMPORT_FORMULA:
                return next((column for column in columns if "formula" in column.casefold()), columns[0])
            return next(
                (column for column in columns if column.casefold() in {"mw", "molecular weight", "molecular_weight"}),
                columns[0],
            )

        source_default = IMPORT_SMILES
        if SDF_MOLECULE_COLUMN not in columns and not any("smiles" in column.casefold() for column in columns):
            source_default = IMPORT_FORMULA if any("formula" in column.casefold() for column in columns) else IMPORT_MW

        window = tk.Toplevel(self)
        window.title("Map imported batch columns")
        window.transient(self)
        window.resizable(False, False)
        window.grab_set()
        body = ttk.Frame(window, padding=14)
        body.pack(fill="both", expand=True)
        id_var = tk.StringVar(value=id_default)
        source_type_var = tk.StringVar(value=source_default)
        source_column_var = tk.StringVar(value=suggested_column(source_default))
        ttk.Label(body, text="ID column").grid(row=0, column=0, sticky="w", padx=(0, 10), pady=4)
        ttk.Combobox(body, textvariable=id_var, values=columns, state="readonly", width=34).grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Label(body, text="Calculation source").grid(row=1, column=0, sticky="w", padx=(0, 10), pady=4)
        source_type_box = ttk.Combobox(
            body,
            textvariable=source_type_var,
            values=(IMPORT_SMILES, IMPORT_FORMULA, IMPORT_MW),
            state="readonly",
            width=34,
        )
        source_type_box.grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Label(body, text="Source column").grid(row=2, column=0, sticky="w", padx=(0, 10), pady=4)
        ttk.Combobox(body, textvariable=source_column_var, values=columns, state="readonly", width=34).grid(
            row=2, column=1, sticky="ew", pady=4
        )
        ttk.Label(
            body,
            text="Only the selected source is used. XinPilot will not compare it with or fall back to another column.",
            wraplength=430,
            justify="left",
        ).grid(row=3, column=0, columnspan=2, sticky="w", pady=(8, 0))
        result: Dict[str, str] = {}

        def source_changed(_event=None):
            source_column_var.set(suggested_column(source_type_var.get()))

        source_type_box.bind("<<ComboboxSelected>>", source_changed)

        def accept():
            result.update({
                "id_column": id_var.get(),
                "source_type": source_type_var.get(),
                "source_column": source_column_var.get(),
            })
            window.destroy()

        controls = ttk.Frame(body)
        controls.grid(row=4, column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(controls, text="Import", command=accept).pack(side="left")
        ttk.Button(controls, text="Cancel", command=window.destroy).pack(side="left", padx=(6, 0))
        window.bind("<Return>", lambda _event: accept())
        window.bind("<Escape>", lambda _event: window.destroy())
        self.wait_window(window)
        return result or None

    def _choose_import_batch_file(self) -> None:
        path_text = filedialog.askopenfilename(
            parent=self,
            title="Import solution-calculator batch",
            filetypes=[
                ("Supported files", "*.sdf *.sd *.xlsx *.xlsm *.csv"),
                ("SDF files", "*.sdf *.sd"),
                ("Excel files", "*.xlsx *.xlsm"),
                ("CSV files", "*.csv"),
            ],
        )
        if not path_text:
            return
        path = Path(path_text)
        try:
            sheet = self._choose_excel_sheet(path) if path.suffix.lower() in {".xlsx", ".xlsm"} else None
            if path.suffix.lower() in {".xlsx", ".xlsm"} and not sheet:
                return
            records, columns = read_batch_import_file(path, sheet)
            mapping = self._choose_import_mapping(columns)
            if mapping is None:
                return
            rows = build_import_batch_rows(
                records,
                id_column=mapping["id_column"],
                source_type=mapping["source_type"],
                source_column=mapping["source_column"],
                target_mm=self.import_common_target_var.get(),
                solvent=self.import_common_solvent_var.get(),
                amount_unit=self.import_common_amount_unit_var.get(),
            )
        except (OSError, ValueError) as exc:
            messagebox.showerror("Could not import batch", str(exc), parent=self)
            return
        self.import_batch_rows[:] = rows
        for row in self.import_batch_rows:
            row["rounding_ul"] = self.import_common_rounding_var.get().strip()
        self.import_file_var.set(f"{path.name} — {len(rows)} row(s), using {mapping['source_type']}")
        if mapping["source_type"] == IMPORT_SMILES:
            self.import_preview_check.configure(state="normal")
        else:
            self.import_preview_var.set(False)
            self.import_preview_check.configure(state="disabled")
            self._close_import_structure_preview(update_checkbox=False)
        self._refresh_import_tree()
        if self.import_batch_rows:
            self.import_tree.selection_set("0")
            self.import_tree.focus("0")
            self.import_tree.see("0")

    def _refresh_import_tree(self) -> None:
        selected = {int(iid) for iid in self.import_tree.selection() if iid.isdigit()}
        self.import_tree.delete(*self.import_tree.get_children())
        for index, row in enumerate(self.import_batch_rows):
            self.import_tree.insert("", "end", iid=str(index), values=(
                row["product_id"], row["identity_source"], row["formula"], row["mw"], row["amount"],
                row["amount_unit"], row["mass_purity"], row["target_mm"], row["solvent"],
                row["final_volume_ul"], row["achieved_mm"], row["status"],
            ))
        retained = [str(index) for index in sorted(selected) if index < len(self.import_batch_rows)]
        if retained:
            self.import_tree.selection_set(retained)
        self._load_selection_editor("import", force=True)

    def _selected_import_indices(self) -> List[int]:
        return sorted(
            int(iid) for iid in self.import_tree.selection()
            if iid.isdigit() and int(iid) < len(self.import_batch_rows)
        )

    def _load_selected_import_row(self, _event=None) -> None:
        self._load_selection_editor("import")
        self._refresh_import_structure_preview()

    def _save_selected_import_row(self) -> bool:
        return self._save_selection_editor("import")

    @staticmethod
    def _calculate_import_row(row: Dict[str, str]) -> bool:
        if row.get("identity_error"):
            row["exact_volume_ul"] = ""
            row["final_volume_ul"] = ""
            row["achieved_mm"] = ""
            row["status"] = row["identity_error"]
            return False
        return SolutionCalculatorApp._calculate_row(row)

    def _calculate_all_import_rows(self) -> None:
        ready = sum(self._calculate_import_row(row) for row in self.import_batch_rows)
        self._refresh_import_tree()
        if not self.import_batch_rows:
            messagebox.showinfo("No imported rows", "Import an SDF, Excel, or CSV file first.", parent=self)
        elif ready < len(self.import_batch_rows):
            messagebox.showwarning(
                "Some rows need input",
                f"Calculated {ready} of {len(self.import_batch_rows)} rows. Review the Status column.",
                parent=self,
            )

    def _apply_import_common_settings(self) -> None:
        for row in self.import_batch_rows:
            row["target_mm"] = self.import_common_target_var.get().strip()
            row["solvent"] = self.import_common_solvent_var.get().strip()
            row["amount_unit"] = self.import_common_amount_unit_var.get()
            row["rounding_ul"] = self.import_common_rounding_var.get().strip()
            self._calculate_import_row(row)
        self._refresh_import_tree()

    def _remove_selected_import_rows(self) -> None:
        for index in reversed(self._selected_import_indices()):
            del self.import_batch_rows[index]
        self._refresh_import_tree()
        self._refresh_import_structure_preview()

    def _send_import_selected_to_single(self) -> None:
        indices = self._selected_import_indices()
        if len(indices) != 1:
            messagebox.showwarning("Select one row", "Select exactly one imported row.", parent=self)
            return
        self._populate_single_from_row(self.import_batch_rows[indices[0]])

    def _toggle_import_structure_preview(self) -> None:
        if not self.import_preview_var.get():
            self._close_import_structure_preview(update_checkbox=False)
            return
        if self._import_preview_window is None or not self._import_preview_window.winfo_exists():
            window = tk.Toplevel(self)
            window.title("Selected imported structure")
            window.transient(self)
            window.geometry("560x500")
            window.minsize(460, 400)
            body = ttk.Frame(window, padding=12)
            body.pack(fill="both", expand=True)
            self._import_preview_title = ttk.Label(body, font=("Segoe UI", 12, "bold"))
            self._import_preview_title.pack(anchor="w", pady=(0, 6))
            self._import_preview_structure = ttk.Label(body, anchor="center")
            self._import_preview_structure.pack(fill="both", expand=True)
            self._import_preview_details = ttk.Label(body, wraplength=520, justify="left")
            self._import_preview_details.pack(fill="x", anchor="w", pady=(8, 0))
            window.protocol("WM_DELETE_WINDOW", self._close_import_structure_preview)
            self._import_preview_window = window
        self._refresh_import_structure_preview()
        self._import_preview_window.lift()

    def _close_import_structure_preview(self, update_checkbox: bool = True) -> None:
        window = self._import_preview_window
        self._import_preview_window = None
        self._import_preview_image = None
        if window is not None and window.winfo_exists():
            window.destroy()
        if update_checkbox and hasattr(self, "import_preview_var"):
            self.import_preview_var.set(False)

    def _refresh_import_structure_preview(self) -> None:
        window = self._import_preview_window
        if not self.import_preview_var.get() or window is None or not window.winfo_exists():
            return
        indices = self._selected_import_indices()
        if len(indices) != 1:
            self._import_preview_title.configure(text="Select one imported row")
            self._import_preview_structure.configure(image="", text="")
            self._import_preview_details.configure(text="The preview remains open and will update when one row is selected.")
            self._import_preview_image = None
            return
        row = self.import_batch_rows[indices[0]]
        smiles = row["identity_value"]
        molecule = Chem.MolFromSmiles(smiles) if Chem is not None and smiles else None
        self._import_preview_title.configure(text=row["product_id"] or "Imported compound")
        if molecule is None or Draw is None:
            self._import_preview_structure.configure(image="", text="Structure unavailable")
            self._import_preview_image = None
        else:
            image = Draw.MolToImage(molecule, size=(500, 340))
            photo = ImageTk.PhotoImage(image)
            self._import_preview_structure.configure(image=photo, text="")
            self._import_preview_image = photo
        self._import_preview_details.configure(
            text=(
                f"Calculation SMILES: {smiles or 'not available'}\n"
                f"Formula: {row['formula'] or 'not available'}    Molecular weight: {row['mw'] or 'not available'}"
            )
        )

    def _export_import_batch_csv(self) -> None:
        self._export_rows_csv(self.import_batch_rows, "Export imported solution calculations")

    def add_products(self, products: Iterable[Dict[str, object]]) -> None:
        for product in products:
            self.batch_rows.append(product_to_batch_row(product))
        self._refresh_batch_tree()
        if self.batch_rows:
            iid = str(len(self.batch_rows) - 1)
            self.batch_tree.selection_set(iid)
            self.batch_tree.focus(iid)
            self.batch_tree.see(iid)

    def _refresh_batch_tree(self) -> None:
        selected_indices = {int(iid) for iid in self.batch_tree.selection() if iid.isdigit()}
        self.batch_tree.delete(*self.batch_tree.get_children())
        for index, row in enumerate(self.batch_rows):
            self.batch_tree.insert("", "end", iid=str(index), values=(
                row["reaction_well"], row["product_name"] or row["product_id"], row["formula"], row["mw"],
                row["amount"], row["amount_unit"], row["mass_purity"], row["target_mm"], row["solvent"],
                row["final_volume_ul"], row["achieved_mm"], row["status"],
            ))
        retained = [str(index) for index in sorted(selected_indices) if index < len(self.batch_rows)]
        if retained:
            self.batch_tree.selection_set(retained)
        self._load_selection_editor("batch", force=True)

    def _selected_batch_indices(self) -> List[int]:
        return sorted(int(iid) for iid in self.batch_tree.selection() if iid.isdigit() and int(iid) < len(self.batch_rows))

    def _load_selected_batch_row(self, _event=None) -> None:
        self._load_selection_editor("batch")
        self._refresh_batch_structure_preview()

    def _toggle_batch_structure_preview(self) -> None:
        if not self.batch_preview_var.get():
            self._close_batch_structure_preview(update_checkbox=False)
            return
        if self._batch_preview_window is None or not self._batch_preview_window.winfo_exists():
            window = tk.Toplevel(self)
            window.title("Selected batch structure")
            window.transient(self)
            window.geometry("560x500")
            window.minsize(460, 400)
            body = ttk.Frame(window, padding=12)
            body.pack(fill="both", expand=True)
            self._batch_preview_title = ttk.Label(body, font=("Segoe UI", 12, "bold"))
            self._batch_preview_title.pack(anchor="w", pady=(0, 6))
            self._batch_preview_structure = ttk.Label(body, anchor="center")
            self._batch_preview_structure.pack(fill="both", expand=True)
            self._batch_preview_details = ttk.Label(body, wraplength=520, justify="left")
            self._batch_preview_details.pack(fill="x", anchor="w", pady=(8, 0))
            window.protocol("WM_DELETE_WINDOW", self._close_batch_structure_preview)
            self._batch_preview_window = window
        self._refresh_batch_structure_preview()
        self._batch_preview_window.lift()

    def _close_batch_structure_preview(self, update_checkbox: bool = True) -> None:
        window = self._batch_preview_window
        self._batch_preview_window = None
        self._batch_preview_image = None
        if window is not None and window.winfo_exists():
            window.destroy()
        if update_checkbox and hasattr(self, "batch_preview_var"):
            self.batch_preview_var.set(False)

    def _refresh_batch_structure_preview(self) -> None:
        window = self._batch_preview_window
        if not self.batch_preview_var.get() or window is None or not window.winfo_exists():
            return
        indices = self._selected_batch_indices()
        if len(indices) != 1:
            self._batch_preview_title.configure(text="Select one batch row")
            self._batch_preview_structure.configure(image="", text="")
            self._batch_preview_details.configure(text="The preview remains open and will update when one row is selected.")
            self._batch_preview_image = None
            return
        row = self.batch_rows[indices[0]]
        name = row["product_name"] or row["product_id"] or row["reaction_well"] or "Selected product"
        smiles = str(row.get("smiles", "") or "").strip()
        molecule = Chem.MolFromSmiles(smiles) if Chem is not None and smiles else None
        self._batch_preview_title.configure(text=name)
        if molecule is None or Draw is None:
            self._batch_preview_structure.configure(image="", text="Structure unavailable")
            self._batch_preview_image = None
        else:
            image = Draw.MolToImage(molecule, size=(500, 340))
            photo = ImageTk.PhotoImage(image)
            self._batch_preview_structure.configure(image=photo, text="")
            self._batch_preview_image = photo
        self._batch_preview_details.configure(
            text=(
                f"Calculation SMILES: {smiles or 'not available'}\n"
                f"Formula: {row['formula'] or 'not available'}    Molecular weight: {row['mw'] or 'not available'}"
            )
        )

    _EDIT_FIELDS = {
        "amount": "amount", "amount_unit": "amount_unit", "mass_purity": "mass_purity",
        "target_mm": "target", "solvent": "solvent", "rounding_ul": "rounding",
    }

    def _init_selection_editors(self):
        self._editor_loading = False
        self._editor_selection = {"batch": None, "import": None}
        self._editor_dirty = {"batch": set(), "import": set()}
        for prefix in ("batch", "import"):
            for field, suffix in self._EDIT_FIELDS.items():
                variable = getattr(self, f"{prefix}_{suffix}_var")
                variable.trace_add("write", lambda *_args, p=prefix, f=field: self._mark_editor_field(p, f))

    def _mark_editor_field(self, prefix, field):
        if not self._editor_loading:
            self._editor_dirty[prefix].add(field)

    def _load_selection_editor(self, prefix, force=False):
        if not hasattr(self, "_editor_selection"):
            return
        indices = tuple(getattr(self, f"_selected_{prefix}_indices")())
        if not force and self._editor_selection[prefix] == indices:
            return
        rows = self.batch_rows if prefix == "batch" else self.import_batch_rows
        self._editor_loading = True
        try:
            for field, suffix in self._EDIT_FIELDS.items():
                values = {str(rows[i].get(field, "")) for i in indices}
                # Blank indicates mixed values; untouched blanks are never applied.
                value = next(iter(values)) if len(values) == 1 else ""
                getattr(self, f"{prefix}_{suffix}_var").set(value)
            self._editor_selection[prefix] = indices
            self._editor_dirty[prefix].clear()
        finally:
            self._editor_loading = False

    def _save_selection_editor(self, prefix):
        indices = getattr(self, f"_selected_{prefix}_indices")()
        if not indices:
            messagebox.showwarning("Select rows", "Select one or more rows to update.", parent=self)
            return False
        if self._editor_selection[prefix] != tuple(indices):
            self._load_selection_editor(prefix)
            messagebox.showinfo("Selection changed", "Enter the changes for the selected rows, then click Update selected.", parent=self)
            return False
        updates = {field: getattr(self, f"{prefix}_{self._EDIT_FIELDS[field]}_var").get().strip()
                   for field in self._editor_dirty[prefix]}
        if not updates:
            return True
        rows = self.batch_rows if prefix == "batch" else self.import_batch_rows
        calculate = self._calculate_row if prefix == "batch" else self._calculate_import_row
        for index in indices:
            rows[index].update(updates)
            calculate(rows[index])
        getattr(self, f"_refresh_{'batch' if prefix == 'batch' else 'import'}_tree")()
        self._load_selection_editor(prefix, force=True)
        return True

    def _save_selected_batch_row(self) -> bool:
        return self._save_selection_editor("batch")

    @staticmethod
    def _calculate_row(row: Dict[str, str]) -> bool:
        try:
            result = prepare_stock(
                amount=row["amount"],
                amount_unit=_amount_unit(row["amount_unit"]),
                target_concentration=row["target_mm"],
                concentration_unit="mM",
                molecular_weight=row["mw"],
                mass_purity_percent=row["mass_purity"],
                rounding_increment_ul=row["rounding_ul"],
            )
        except ValueError as exc:
            row["exact_volume_ul"] = ""
            row["final_volume_ul"] = ""
            row["achieved_mm"] = ""
            row["status"] = str(exc)
            return False
        row["exact_volume_ul"] = format_number(result.exact_volume_ul)
        row["final_volume_ul"] = format_number(result.final_volume_ul)
        row["achieved_mm"] = format_number(result.achieved_concentration_mm)
        row["status"] = "Ready"
        return True

    def _calculate_all_rows(self) -> None:
        ready = sum(self._calculate_row(row) for row in self.batch_rows)
        self._refresh_batch_tree()
        if not self.batch_rows:
            messagebox.showinfo("No batch rows", "Send products from XinPilot or add a blank row first.", parent=self)
        elif ready < len(self.batch_rows):
            messagebox.showwarning(
                "Some rows need input",
                f"Calculated {ready} of {len(self.batch_rows)} rows. Review the Status column for missing or invalid values.",
                parent=self,
            )

    def _apply_common_settings(self) -> None:
        for row in self.batch_rows:
            row["target_mm"] = self.batch_common_target_var.get().strip()
            row["solvent"] = self.batch_common_solvent_var.get().strip()
            row["rounding_ul"] = self.batch_common_rounding_var.get().strip()
            self._calculate_row(row)
        self._refresh_batch_tree()

    def _add_blank_batch_row(self) -> None:
        self.add_products([{}])

    def _remove_selected_rows(self) -> None:
        indices = self._selected_batch_indices()
        if not indices:
            return
        for index in reversed(indices):
            del self.batch_rows[index]
        self._refresh_batch_tree()
        self._refresh_batch_structure_preview()

    def _send_selected_to_single(self) -> None:
        indices = self._selected_batch_indices()
        if len(indices) != 1:
            messagebox.showwarning("Select one row", "Select exactly one batch row.", parent=self)
            return
        self._populate_single_from_row(self.batch_rows[indices[0]])

    def _populate_single_from_row(self, row: Dict[str, str]) -> None:
        self.single_name_var.set(row["product_name"] or row["product_id"])
        self.single_smiles_var.set(row["smiles"])
        self.single_formula_var.set(row["formula"])
        self.single_mw_var.set(row["mw"])
        self.single_exact_mass_var.set(row["exact_mass"])
        self.single_solvent_var.set(row["solvent"])
        self.single_amount_var.set(row["amount"])
        self.single_amount_unit_var.set(row["amount_unit"])
        self.single_purity_var.set(row["mass_purity"])
        self.single_concentration_var.set(row["target_mm"])
        self.single_concentration_unit_var.set("mM")
        self.single_rounding_var.set(row["rounding_ul"])
        self.mode_var.set(MODE_PREPARE)
        self._update_single_mode_hint()
        self.notebook.select(self.single_tab)

    @staticmethod
    def _batch_recipe(row: Dict[str, str]) -> str:
        name = row["product_name"] or row["product_id"] or row["reaction_well"] or "Compound"
        if row["status"] != "Ready":
            return f"{name}: not calculated ({row['status']})."
        return (
            f"{name} (RXN {row['reaction_well'] or 'N/A'}): recovered {row['amount']} {row['amount_unit']}; "
            f"dissolve in {row['solvent'] or 'selected solvent'}, then make up to {row['final_volume_ul']} µL final volume. "
            f"Achieved concentration {row['achieved_mm']} mM."
        )

    def _copy_batch_notes(self) -> None:
        notes = "\n".join(self._batch_recipe(row) for row in self.batch_rows)
        if notes:
            self.clipboard_clear()
            self.clipboard_append(notes)

    def _export_batch_csv(self) -> None:
        self._export_rows_csv(self.batch_rows, "Export solution calculations")

    def _export_rows_csv(self, rows: List[Dict[str, str]], title: str) -> None:
        if not rows:
            messagebox.showinfo("No batch rows", "There are no rows to export.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self,
            title=title,
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv")],
        )
        if not path:
            return
        fields = (
            "reaction_well", "lcms_well", "reaction_name", "product_id", "product_name", "smiles", "formula",
            "exact_mass", "mw", "amount", "amount_unit", "mass_purity", "target_mm", "solvent", "rounding_ul",
            "exact_volume_ul", "final_volume_ul", "achieved_mm", "uv_purity_reference", "uv_purity_source",
            "lcms_reference_smiles", "identity_source", "identity_value", "status",
        )
        try:
            with open(path, "w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
                writer.writeheader()
                writer.writerows(rows)
        except OSError as exc:
            messagebox.showerror("Could not export CSV", str(exc), parent=self)
            return
        messagebox.showinfo("CSV exported", f"Saved {len(rows)} row(s) to:\n{path}", parent=self)

    @staticmethod
    def _set_text(widget: tk.Text, value: str) -> None:
        widget.configure(state="normal")
        widget.delete("1.0", "end")
        widget.insert("1.0", value)
        widget.configure(state="disabled")


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Solution and stock calculator")
    parser.add_argument("--input", help="Temporary XinPilot product-transfer JSON file")
    args = parser.parse_args(argv)
    payload: Optional[Dict[str, object]] = None
    if args.input:
        try:
            payload = load_xinpilot_payload(args.input)
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            root = tk.Tk()
            root.withdraw()
            messagebox.showerror("Could not load XinPilot products", str(exc), parent=root)
            root.destroy()
    app = SolutionCalculatorApp(payload)
    app.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
